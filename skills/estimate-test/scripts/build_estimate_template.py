#!/usr/bin/env python3
"""Sinh templates/template_test_estimate.xlsx cho skill estimate-test.

Chạy lại bất cứ lúc nào để regen template sau khi sửa estimate_template_data.py.
Layout (row constant) được khai báo tường minh ở RC_* vì formula ở các sheet khác
tham chiếu tuyệt đối vào Rate Card.
"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
import estimate_template_data as D  # noqa: E402

def _templates_dir() -> Path:
    """Dò ngược lên tìm thư mục chứa `templates/`.

    Hỗ trợ cả 2 layout: bản portable `~/.claude/skills/estimate-test/` (templates nằm
    cạnh scripts/) và bản trong project `<workspace>/skills/estimate-test/` (templates
    nằm ở workspace root). Nhờ vậy 2 bản script giữ nội dung giống nhau.
    """
    start = Path(__file__).resolve().parent
    for p in (start, *start.parents):
        if (p / "templates").is_dir():
            return p / "templates"
    return start.parent / "templates"


OUT_PATH = _templates_dir() / "template_test_estimate.xlsx"

RC = "'Rate Card'"

# Layout Rate Card suy ra từ độ dài các bảng dữ liệu — thêm/bớt activity là mọi
# section tự dịch, không phải sửa row constant bằng tay (mỗi section: title, header,
# data..., 1 dòng trắng).
RC_RATE_SEC, RC_RATE_HDR, RC_RATE_FIRST = 3, 4, 5
RC_RATE_LAST = RC_RATE_FIRST + len(D.BASE_RATES) - 1
RC_LEVEL_SEC = RC_RATE_LAST + 2
RC_LEVEL_FIRST = RC_LEVEL_SEC + 2
RC_LEVEL_LAST = RC_LEVEL_FIRST + len(D.LEVELS) - 1
RC_MIN_SEC = RC_LEVEL_LAST + 2
RC_MIN_FIRST = RC_MIN_SEC + 2
RC_MIN_LAST = RC_MIN_FIRST + len(D.MIN_LEVELS) - 1
RC_MULT_SEC = RC_MIN_LAST + 2
RC_MULT_FIRST = RC_MULT_SEC + 2
RC_MULT_LAST = RC_MULT_FIRST + len(D.MULTIPLIERS) - 1
RC_PARAM_SEC = RC_MULT_LAST + 2
RC_PARAM_FIRST = RC_PARAM_SEC + 2
RC_PARAM_LAST = RC_PARAM_FIRST + len(D.PARAMS) - 1
RC_CRIT_SEC = RC_PARAM_LAST + 2
RC_CRIT_FIRST = RC_CRIT_SEC + 2

# Dòng TOTAL phải nằm NGOÀI range mà nó SUM, nếu không Excel báo circular reference
# (SUMIF của Estimate Summary cũng quét đúng range DETAIL_FIRST..DETAIL_SUM_LAST này).
DETAIL_FIRST, DETAIL_SUM_LAST = 4, 200
DETAIL_SAMPLE_LAST = DETAIL_FIRST + len(D.SAMPLE_ROWS) - 1
DETAIL_FORMULA_LAST = DETAIL_SUM_LAST
DETAIL_TOTAL_ROW = DETAIL_SUM_LAST + 2

TITLE_FILL = PatternFill("solid", fgColor="1F3864")
SECTION_FILL = PatternFill("solid", fgColor="D9E2F3")
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
CALC_FILL = PatternFill("solid", fgColor="F2F2F2")
TOTAL_FILL = PatternFill("solid", fgColor="FCE4D6")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def put_title(ws, text, span):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=13, color="FFFFFF")
    ws["A1"].fill = TITLE_FILL
    ws.merge_cells(f"A1:{get_column_letter(span)}1")
    ws.row_dimensions[1].height = 24


def put_section(ws, row, text, span):
    ws.cell(row=row, column=1, value=text).font = Font(bold=True, size=11, color="1F3864")
    for col in range(1, span + 1):
        ws.cell(row=row, column=col).fill = SECTION_FILL


def put_header(ws, row, headers):
    for idx, name in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=idx, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = BORDER


def put_row(ws, row, values, fills=None, formats=None):
    for idx, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=idx, value=value)
        cell.border = BORDER
        cell.alignment = WRAP
        if fills and idx in fills:
            cell.fill = fills[idx]
        if formats and idx in formats:
            cell.number_format = formats[idx]


def set_widths(ws, widths):
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def build_rate_card(ws):
    put_title(ws, "TEST ESTIMATE — RATE CARD (chỉ cần tune ở sheet này)", 7)
    set_widths(ws, [26, 46, 18, 12, 12, 12, 12])
    rate_fmt = {4: "0.00", 5: "0.00", 6: "0.00", 7: "0.00"}

    put_section(ws, RC_RATE_SEC, "1. BASE RATE — giờ / đơn vị, tính cho level Middle (baseline). "
                                 "G* = chỉ định AI gen · A* = người làm · S* = support · R1 = rework", 7)
    put_header(ws, RC_RATE_HDR, ["Code", "Activity", "Đơn vị khối lượng", "S", "M", "L", "XL"])
    for offset, item in enumerate(D.BASE_RATES):
        put_row(ws, RC_RATE_FIRST + offset, list(item),
                fills={4: INPUT_FILL, 5: INPUT_FILL, 6: INPUT_FILL, 7: INPUT_FILL}, formats=rate_fmt)

    put_section(ws, RC_LEVEL_SEC, "2. LEVEL FACTOR — nhân vào base rate theo level người thực hiện", 7)
    put_header(ws, RC_LEVEL_SEC + 1, ["Level", "Factor", "Rank", "Ghi chú"])
    for offset, item in enumerate(D.LEVELS):
        put_row(ws, RC_LEVEL_FIRST + offset, list(item), fills={2: INPUT_FILL}, formats={2: "0.00", 3: "0"})

    put_section(ws, RC_MIN_SEC, "3. LEVEL TỐI THIỂU THEO ACTIVITY — gán thấp hơn thì cộng phụ phí senior review. "
                               "Nhạy level = 0 → Level factor cứng 1.00 (nhóm G*: giờ chỉ định AI gen không đổi theo level)", 7)
    put_header(ws, RC_MIN_SEC + 1, ["Code", "Activity", "Level tối thiểu", "Min rank",
                                    "Phụ phí senior review", "Nhạy level? (1/0)"])
    for offset, item in enumerate(D.MIN_LEVELS):
        put_row(ws, RC_MIN_FIRST + offset, list(item), fills={5: INPUT_FILL, 6: INPUT_FILL},
                formats={4: "0", 5: "0%", 6: "0"})

    put_section(ws, RC_MULT_SEC, "4. MULTIPLIER — nhập tay vào cột Multiplier của Estimate Detail (nhân dồn nếu nhiều yếu tố)", 7)
    put_header(ws, RC_MULT_SEC + 1, ["Yếu tố", "Giá trị", "Ghi chú"])
    for offset, item in enumerate(D.MULTIPLIERS):
        put_row(ws, RC_MULT_FIRST + offset, list(item), fills={2: INPUT_FILL}, formats={2: "0.00"})

    put_section(ws, RC_PARAM_SEC, "5. THAM SỐ PHÁI SINH", 7)
    put_header(ws, RC_PARAM_SEC + 1, ["Tham số", "Giá trị", "Ghi chú"])
    for offset, (name, value, fmt, note) in enumerate(D.PARAMS):
        put_row(ws, RC_PARAM_FIRST + offset, [name, value, note], fills={2: INPUT_FILL}, formats={2: fmt})

    put_section(ws, RC_CRIT_SEC, "6. TIÊU CHÍ CHẤM ĐỘ PHỨC TẠP TASK (S / M / L / XL)", 7)
    put_header(ws, RC_CRIT_SEC + 1, ["Tiêu chí", "S", "M", "L", "XL"])
    for offset, item in enumerate(D.COMPLEXITY_CRITERIA):
        put_row(ws, RC_CRIT_FIRST + offset, list(item), fills={1: CALC_FILL})


def add_defined_names(wb):
    for offset, (name, _value, _fmt, _note) in enumerate(D.PARAMS):
        ref = f"{RC}!$B${RC_PARAM_FIRST + offset}"
        wb.defined_names.add(DefinedName(name, attr_text=ref))


def detail_formulas(row):
    rate_tbl = f"{RC}!$A${RC_RATE_FIRST}:$G${RC_RATE_LAST}"
    name_tbl = f"{RC}!$A${RC_RATE_FIRST}:$C${RC_RATE_LAST}"
    lvl_tbl = f"{RC}!$A${RC_LEVEL_FIRST}:$D${RC_LEVEL_LAST}"
    min_tbl = f"{RC}!$A${RC_MIN_FIRST}:$F${RC_MIN_LAST}"
    hdr = f"{RC}!$D${RC_RATE_HDR}:$G${RC_RATE_HDR}"
    return {
        "C": f'=IFERROR(VLOOKUP($B{row},{name_tbl},2,FALSE),"")',
        "F": f'=IFERROR(VLOOKUP($B{row},{name_tbl},3,FALSE),"")',
        "G": f'=IFERROR(VLOOKUP($B{row},{rate_tbl},MATCH($D{row},{hdr},0)+3,FALSE),"")',
        # Cột F bảng 3 = 0 → activity KHÔNG nhạy level (nhóm G*: chỉ định AI gen,
        # giờ chạy skill + chờ output như nhau ở mọi level) → factor cứng 1.00.
        "I": (f'=IFERROR(IF(VLOOKUP($B{row},{min_tbl},6,FALSE)=0,1,'
              f'VLOOKUP($H{row},{lvl_tbl},2,FALSE)),"")'),
        "K": (f'=IFERROR(IF(VLOOKUP($H{row},{lvl_tbl},3,FALSE)<VLOOKUP($B{row},{min_tbl},4,FALSE),'
              f'1+VLOOKUP($B{row},{min_tbl},5,FALSE),1),"")'),
        "L": f'=IFERROR($E{row}*$G{row}*$I{row}*$J{row}*$K{row},"")',
    }


def build_detail(ws):
    put_title(ws, "ESTIMATE DETAIL — 1 dòng = 1 activity của 1 function (ô vàng nhập tay, ô xám tự tính)", 13)
    set_widths(ws, [22, 9, 40, 12, 10, 18, 11, 12, 12, 11, 12, 12, 40])
    put_header(ws, 3, ["Function", "Activity", "Activity name", "Complexity", "Volume", "Đơn vị",
                       "Base rate (h)", "Level", "Level factor", "Multiplier", "Level check",
                       "Adjusted (h)", "Note"])
    ws.freeze_panes = "A4"

    input_fills = {1: INPUT_FILL, 2: INPUT_FILL, 4: INPUT_FILL, 5: INPUT_FILL, 8: INPUT_FILL,
                   10: INPUT_FILL, 13: INPUT_FILL}
    calc_fills = {3: CALC_FILL, 6: CALC_FILL, 7: CALC_FILL, 9: CALC_FILL, 11: CALC_FILL, 12: CALC_FILL}
    fills = {**input_fills, **calc_fills}
    formats = {5: "0.00", 7: "0.00", 9: "0.00", 10: "0.00", 11: "0.00", 12: "0.00"}

    for row in range(DETAIL_FIRST, DETAIL_FORMULA_LAST + 1):
        sample = D.SAMPLE_ROWS[row - DETAIL_FIRST] if row <= DETAIL_SAMPLE_LAST else None
        f = detail_formulas(row)
        values = [
            sample[0] if sample else None, sample[1] if sample else None, f["C"],
            sample[2] if sample else None, sample[3] if sample else None, f["F"], f["G"],
            sample[4] if sample else None, f["I"], sample[5] if sample else None, f["K"], f["L"],
            sample[6] if sample else None,
        ]
        put_row(ws, row, values, fills=fills, formats=formats)

    put_row(ws, DETAIL_TOTAL_ROW,
            [None] * 11 + [f"=SUM(L{DETAIL_FIRST}:L{DETAIL_SUM_LAST})", "Tổng chưa buffer"],
            fills={12: TOTAL_FILL}, formats={12: "0.00"})
    ws.cell(row=DETAIL_TOTAL_ROW, column=12).font = Font(bold=True)

    dv_code = DataValidation(type="list", formula1='"%s"' % ",".join(r[0] for r in D.BASE_RATES), allow_blank=True)
    dv_cplx = DataValidation(type="list", formula1='"S,M,L,XL"', allow_blank=True)
    dv_level = DataValidation(type="list", formula1='"%s"' % ",".join(r[0] for r in D.LEVELS), allow_blank=True)
    for dv, col in ((dv_code, "B"), (dv_cplx, "D"), (dv_level, "H")):
        ws.add_data_validation(dv)
        dv.add(f"{col}{DETAIL_FIRST}:{col}{DETAIL_SUM_LAST}")


def build_summary(ws):
    put_title(ws, "ESTIMATE SUMMARY — effort theo hoạt động, theo level, theo function", 5)
    set_widths(ws, [26, 46, 16, 18, 20])
    dl_code = f"'Estimate Detail'!$B${DETAIL_FIRST}:$B${DETAIL_SUM_LAST}"
    dl_level = f"'Estimate Detail'!$H${DETAIL_FIRST}:$H${DETAIL_SUM_LAST}"
    dl_func = f"'Estimate Detail'!$A${DETAIL_FIRST}:$A${DETAIL_SUM_LAST}"
    dl_hours = f"'Estimate Detail'!$L${DETAIL_FIRST}:$L${DETAIL_SUM_LAST}"

    put_section(ws, 3, "1. EFFORT THEO HOẠT ĐỘNG", 5)
    put_header(ws, 4, ["Code", "Activity", "Manhour", "% tổng"])
    first, last = 5, 5 + len(D.BASE_RATES) - 1
    for offset, item in enumerate(D.BASE_RATES):
        row = first + offset
        put_row(ws, row, [item[0], item[1], f'=SUMIF({dl_code},$A{row},{dl_hours})',
                          f'=IFERROR(C{row}/$C${last + 1},"")'],
                fills={3: CALC_FILL, 4: CALC_FILL}, formats={3: "0.00", 4: "0%"})
    sub, buf, tot_h, tot_d = last + 1, last + 2, last + 3, last + 4
    put_row(ws, sub, [None, "Subtotal (chưa buffer)", f"=SUM(C{first}:C{last})"],
            fills={3: TOTAL_FILL}, formats={3: "0.00"})
    put_row(ws, buf, [None, "Buffer risk", f"=C{sub}*buffer_pct"], fills={3: TOTAL_FILL}, formats={3: "0.00"})
    put_row(ws, tot_h, [None, "TOTAL manhour", f"=C{sub}+C{buf}"], fills={3: TOTAL_FILL}, formats={3: "0.00"})
    put_row(ws, tot_d, [None, "TOTAL manday", f"=C{tot_h}/hours_per_manday"],
            fills={3: TOTAL_FILL}, formats={3: "0.00"})
    for row in (sub, buf, tot_h, tot_d):
        ws.cell(row=row, column=2).font = Font(bold=True)
        ws.cell(row=row, column=3).font = Font(bold=True)

    lvl_sec = tot_d + 2
    put_section(ws, lvl_sec, "2. EFFORT THEO LEVEL NHÂN SỰ (không tính số người — chỉ level)", 5)
    put_header(ws, lvl_sec + 1, ["Level", "Manhour", "Manday (chưa buffer)", "Manday (có buffer)"])
    lvl_first = lvl_sec + 2
    for offset, item in enumerate(D.LEVELS):
        row = lvl_first + offset
        put_row(ws, row, [item[0], f'=SUMIF({dl_level},$A{row},{dl_hours})',
                          f"=B{row}/hours_per_manday", f"=B{row}*(1+buffer_pct)/hours_per_manday"],
                fills={2: CALC_FILL, 3: CALC_FILL, 4: CALC_FILL}, formats={2: "0.00", 3: "0.00", 4: "0.00"})
    lvl_last = lvl_first + len(D.LEVELS) - 1
    put_row(ws, lvl_last + 1, ["Tổng", f"=SUM(B{lvl_first}:B{lvl_last})",
                               f"=SUM(C{lvl_first}:C{lvl_last})", f"=SUM(D{lvl_first}:D{lvl_last})"],
            fills={2: TOTAL_FILL, 3: TOTAL_FILL, 4: TOTAL_FILL},
            formats={2: "0.00", 3: "0.00", 4: "0.00"})
    ws.cell(row=lvl_last + 1, column=1).font = Font(bold=True)

    fn_sec = lvl_last + 3
    put_section(ws, fn_sec, "3. EFFORT THEO FUNCTION / MODULE", 5)
    put_header(ws, fn_sec + 1, ["Function", "Manhour", "Manday (có buffer)"])
    fn_first = fn_sec + 2
    functions = list(dict.fromkeys(r[0] for r in D.SAMPLE_ROWS))  # giữ thứ tự xuất hiện
    for row in range(fn_first, fn_first + 12):
        idx = row - fn_first
        value = functions[idx] if idx < len(functions) else None
        put_row(ws, row, [value, f'=IFERROR(SUMIF({dl_func},$A{row},{dl_hours}),"")',
                          f'=IFERROR(B{row}*(1+buffer_pct)/hours_per_manday,"")'],
                fills={1: INPUT_FILL, 2: CALC_FILL, 3: CALC_FILL}, formats={2: "0.00", 3: "0.00"})

    note_row = fn_first + 13
    ws.cell(row=note_row, column=1,
            value=("Ghi chú: bản estimate này KHÔNG có cột số nhân sự / duration — giai đoạn estimate "
                   "chỉ chốt độ phức tạp, khối lượng và LEVEL người thực hiện. Phân bổ người và lịch "
                   "làm ở bước planning."))
    ws.cell(row=note_row, column=1).alignment = WRAP
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=5)
    ws.row_dimensions[note_row].height = 30


def build_assumptions(ws):
    put_title(ws, "ASSUMPTIONS & RISKS — mọi giả định làm nên con số estimate", 6)
    set_widths(ws, [6, 16, 58, 40, 14, 30])
    put_header(ws, 3, ["#", "Loại", "Nội dung", "Ảnh hưởng tới estimate", "Độ tin cậy", "Nguồn"])
    for offset, item in enumerate(D.ASSUMPTION_ROWS):
        put_row(ws, 4 + offset, [offset + 1, *item], fills={1: CALC_FILL})
    for row in range(4 + len(D.ASSUMPTION_ROWS), 4 + len(D.ASSUMPTION_ROWS) + 12):
        put_row(ws, row, [row - 3, None, None, None, None, None], fills={1: CALC_FILL})
    dv_type = DataValidation(type="list", formula1='"Assumption,Open point,Risk,Level warning"', allow_blank=True)
    ws.add_data_validation(dv_type)
    dv_type.add("B4:B100")
    dv_conf = DataValidation(type="list", formula1='"Cao,Trung bình,Thấp"', allow_blank=True)
    ws.add_data_validation(dv_conf)
    dv_conf.add("E4:E100")


def main():
    wb = Workbook()
    build_rate_card(wb.active)
    wb.active.title = "Rate Card"
    add_defined_names(wb)
    build_detail(wb.create_sheet("Estimate Detail"))
    build_summary(wb.create_sheet("Estimate Summary"))
    build_assumptions(wb.create_sheet("Assumptions & Risks"))
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_PATH)
    print(f"OK -> {OUT_PATH}")


if __name__ == "__main__":
    main()
