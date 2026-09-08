"""Test cho scripts/write_testcase_xlsx.py va verify_testcase_xlsx.py."""
import csv
import sys
from pathlib import Path

import openpyxl
import pytest

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
TEMPLATE = Path(__file__).resolve().parents[1] / "templates" / "template_testcase.xlsx"
sys.path.insert(0, str(SCRIPTS))

import verify_testcase_xlsx as V  # noqa: E402
from write_testcase_xlsx import CSV_COLS, ghi  # noqa: E402


def _csv(tmp_path, rows):
    p = tmp_path / "tc.csv"
    with open(p, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=CSV_COLS)
        w.writeheader()
        w.writerows(rows)
    return p


def _tc(i, **kw):
    d = {"Classification I": "Common Test Cases", "Classification II": "Navigation",
         "Classification III": "tab", "Test subject": f"case {i}", "Priority": "High",
         "Test case ID": f"TC_{i:03d}", "Pre-condition": "1. logged in",
         "Steps to reproduce": "1. Open screen\n2. Click tab",
         "Test data": "-", "Expected result": "1. Screen shown\n2.1. Tab opened\n2.2. Area shown"}
    d.update(kw)
    return d


def test_ghi_xoa_sach_ket_qua_thi_hanh_mau_cua_template(tmp_path):
    """Template co san 'Passed'/'HuongDT'/2023-07-07 o K..T — bàn giao kem la pass gia."""
    # Arrange
    src = _csv(tmp_path, [_tc(1), _tc(2)])
    out = tmp_path / "o.xlsx"

    # Act
    ghi(src, out, TEMPLATE, sheet_name="m")

    # Assert
    ws = openpyxl.load_workbook(out)["m"]
    sot = [c for r in range(12, 14) for c in "KLMNOPQRST" if ws[f"{c}{r}"].value is not None]
    assert sot == []


def test_ghi_chen_them_dong_khi_vuot_9_dong_co_san(tmp_path):
    """Template chi ke san 9 dong (12-20)."""
    # Arrange
    src = _csv(tmp_path, [_tc(i) for i in range(1, 26)])
    out = tmp_path / "o.xlsx"

    # Act
    n, _, _, _ = ghi(src, out, TEMPLATE, sheet_name="m")

    # Assert
    ws = openpyxl.load_workbook(out)["m"]
    assert n == 25
    assert ws["F36"].value == "TC_025"        # 12 + 25 - 1
    assert ws["F37"].value is None


def test_ghi_gom_nhom_cot_phan_loai_bang_merge_doc(tmp_path):
    # Arrange — 3 case cung Classification I
    src = _csv(tmp_path, [_tc(i) for i in range(1, 4)])
    out = tmp_path / "o.xlsx"

    # Act
    ghi(src, out, TEMPLATE, sheet_name="m")

    # Assert — A12:A14 merge lai, chi o dau giu gia tri
    ws = openpyxl.load_workbook(out)["m"]
    assert "A12:A14" in [str(m) for m in ws.merged_cells.ranges]
    assert ws["A12"].value == "Common Test Cases"


def test_ghi_khong_de_dong_nao_bi_cat_chu(tmp_path):
    # Arrange — noi dung dai
    src = _csv(tmp_path, [_tc(1, **{"Expected result": "\n".join(f"{i}. ket qua rat dai " * 6 for i in range(1, 4))})])
    out = tmp_path / "o.xlsx"

    # Act
    ghi(src, out, TEMPLATE, sheet_name="m")

    # Assert
    ws = openpyxl.load_workbook(out)["m"]
    assert ws.row_dimensions[12].height and ws.row_dimensions[12].height > 30


def test_ghi_dat_ten_sheet_theo_chuc_nang(tmp_path):
    # Arrange
    src = _csv(tmp_path, [_tc(1)])
    out = tmp_path / "o.xlsx"

    # Act
    _, _, ten, _ = ghi(src, out, TEMPLATE, sheet_name="event_edit_schedule")

    # Assert
    assert ten == "event_edit_schedule"
    assert "Function {Name}" not in openpyxl.load_workbook(out).sheetnames


def test_ghi_cat_ten_sheet_qua_31_ky_tu(tmp_path):
    # Arrange
    src = _csv(tmp_path, [_tc(1)])
    out = tmp_path / "o.xlsx"

    # Act
    _, _, ten, _ = ghi(src, out, TEMPLATE, sheet_name="x" * 40)

    # Assert — Excel gioi han 31 ky tu
    assert len(ten) == 31


def test_ghi_dung_han_khi_csv_thieu_cot(tmp_path):
    # Arrange
    p = tmp_path / "xau.csv"
    p.write_text("Test case ID,Title\nTC_001,abc\n", encoding="utf-8")

    # Act + Assert
    with pytest.raises(SystemExit):
        ghi(p, tmp_path / "o.xlsx", TEMPLATE)


def test_verify_bat_ket_qua_gia_con_sot(tmp_path):
    # Arrange
    src = _csv(tmp_path, [_tc(1)])
    out = tmp_path / "o.xlsx"
    ghi(src, out, TEMPLATE, sheet_name="m")
    wb = openpyxl.load_workbook(out)
    wb["m"]["K12"] = "Passed"
    wb.save(out)

    # Act
    wb2 = openpyxl.load_workbook(out)
    dat, ct = V.check_khong_ket_qua_gia(wb2["m"], 1, None)

    # Assert
    assert dat is False and "K12" in ct


def test_verify_chap_nhan_expected_danh_so_phan_cap(tmp_path):
    """1 buoc duoc phep co nhieu diem verify: 2 -> 2.1 / 2.2."""
    # Arrange
    src = _csv(tmp_path, [_tc(1)])
    out = tmp_path / "o.xlsx"
    ghi(src, out, TEMPLATE, sheet_name="m")

    # Act
    dat, _ = V.check_steps_khop_expected(openpyxl.load_workbook(out)["m"], 1, None)

    # Assert
    assert dat is True


def test_verify_bat_buoc_thieu_expected():
    """Buoc 3 khong co expected nao -> phai FAIL."""
    # Arrange
    wb = openpyxl.Workbook(); ws = wb.active
    ws["H12"] = "1. a\n2. b\n3. c"
    ws["J12"] = "1. x\n2. y"

    # Act
    dat, ct = V.check_steps_khop_expected(ws, 1, None)

    # Assert
    assert dat is False and "[3]" in ct


def test_verify_bat_test_case_id_trung():
    # Arrange
    wb = openpyxl.Workbook(); ws = wb.active
    ws["F12"], ws["F13"] = "TC_001", "TC_001"

    # Act
    dat, ct = V.check_tcid_duy_nhat(ws, 2, None)

    # Assert
    assert dat is False and "TC_001" in ct


# ---------- danh dau diem chua chot (dong bo voi /gen-test-plan) ----------

def test_to_do_red_bat_ky_to_ca_token_giua_cau(tmp_path):
    """Khac /gen-test-plan: trong test case, token o dau cung la loi -> deu to."""
    # Arrange
    from xlsx_row_ops import DO, to_do_red
    src = _csv(tmp_path, [_tc(1, **{"Expected result": "1. Response time Cần xác nhận"})])
    out = tmp_path / "o.xlsx"

    # Act
    ghi(src, out, TEMPLATE, sheet_name="m")

    # Assert
    ws = openpyxl.load_workbook(out)["m"]
    assert ws["J12"].font.color.rgb == DO


def test_verify_fail_khi_expected_result_con_chua_chot(tmp_path):
    """Test case khong co ket qua mong doi chot thi khong chay duoc -> cam ban giao."""
    # Arrange
    src = _csv(tmp_path, [_tc(1, **{"Expected result": "1. Ngưỡng Cần xác nhận"})])
    out = tmp_path / "o.xlsx"
    ghi(src, out, TEMPLATE, sheet_name="m")

    # Act
    dat, ct = V.check_expected_khong_chua_chot(openpyxl.load_workbook(out)["m"], 1, None)

    # Assert
    assert dat is False and "J12" in ct


def test_verify_fail_khi_o_chua_chot_chua_to_do(tmp_path):
    # Arrange — ghi xong roi co tinh xoa mau di
    from openpyxl.styles import Font
    src = _csv(tmp_path, [_tc(1, **{"Pre-condition": "1. Tài khoản Cần xác nhận"})])
    out = tmp_path / "o.xlsx"
    ghi(src, out, TEMPLATE, sheet_name="m")
    wb = openpyxl.load_workbook(out)
    wb["m"]["G12"].font = Font(color="FF000000")
    wb.save(out)

    # Act
    dat, ct = V.check_chua_chot_da_to_do(openpyxl.load_workbook(out)["m"], 1, None)

    # Assert
    assert dat is False and "G12" in ct


def test_verify_pass_khi_khong_co_token_nao(tmp_path):
    # Arrange
    src = _csv(tmp_path, [_tc(1), _tc(2)])
    out = tmp_path / "o.xlsx"
    ghi(src, out, TEMPLATE, sheet_name="m")

    # Act
    ws = openpyxl.load_workbook(out)["m"]

    # Assert
    assert V.check_expected_khong_chua_chot(ws, 2, None)[0] is True
    assert V.check_chua_chot_da_to_do(ws, 2, None)[0] is True


# --- Cover + ToC + header sheet test case (write_meta_cells.py) ---

META_DU = {"project_name": "Gettii Lite (GTL)", "creator": "HuongDT",
           "user_story": "GTL-1234", "purpose": "This document is used to verify the menu",
           "test_environment": "macOS 14, Chrome 127", "reviewer": "QuyenNT",
           "review_date": "2026-09-04", "reference": "外部設計書 v1.5",
           "module_description": "Buyer navigator menu of the lottery flow",
           "screen_name": "マイチケット"}


def test_ghi_cover_voi_meta_du_thi_khong_con_placeholder_template(tmp_path):
    """14 o Cover cua template deu la placeholder `<...>` — phai bi ghi de hoan toan."""
    # Arrange
    src = _csv(tmp_path, [_tc(1)])
    out = tmp_path / "o.xlsx"

    # Act
    ghi(src, out, TEMPLATE, sheet_name="buyer_navigator_menu", meta=META_DU)

    # Assert
    cv = openpyxl.load_workbook(out)["Cover"]
    sot = [c.coordinate for r in cv.iter_rows() for c in r
           if isinstance(c.value, str) and V.PLACEHOLDER.search(c.value)]
    assert sot == []
    assert cv["C3"].value == "Gettii Lite (GTL)"
    assert cv["C4"].value == "buyer_navigator_menu"      # module_name suy tu ten sheet
    assert cv["C11"].value == "1.0"                      # version mac dinh


def test_ghi_cover_thieu_meta_thi_ghi_chua_chot_va_to_chu_do(tmp_path):
    """De trong thi review khong phan biet duoc 'chua co tin' voi 'khong can dien'."""
    # Arrange
    from xlsx_row_ops import DO, TOKEN_CHUA_CHOT
    src = _csv(tmp_path, [_tc(1)])
    out = tmp_path / "o.xlsx"

    # Act
    _, _, _, chua_chot = ghi(src, out, TEMPLATE, sheet_name="m", meta=None)

    # Assert
    cv = openpyxl.load_workbook(out)["Cover"]
    assert cv["C3"].value == TOKEN_CHUA_CHOT             # project_name khong suy ra duoc
    assert cv["C3"].font.color.rgb == DO
    assert chua_chot > 0


def test_ghi_toc_dong_7_tro_dung_ten_sheet_test_case(tmp_path):
    """ToC B7 cua template la 'Function {Name}' — phai doi theo ten sheet that."""
    # Arrange
    src = _csv(tmp_path, [_tc(1)])
    out = tmp_path / "o.xlsx"

    # Act
    ghi(src, out, TEMPLATE, sheet_name="event_edit_schedule", meta=META_DU)

    # Assert
    toc = openpyxl.load_workbook(out)["Table of content"]
    assert toc["B7"].value == "event_edit_schedule"
    assert "event_edit_schedule" in toc["C7"].value
    assert META_DU["module_description"] in toc["C7"].value


def test_verify_fail_khi_cover_con_placeholder_cua_template(tmp_path):
    """Regression: truoc day verify chi soi sheet test case nen Cover trang van pass."""
    # Arrange
    src = _csv(tmp_path, [_tc(1)])
    out = tmp_path / "o.xlsx"
    ghi(src, out, TEMPLATE, sheet_name="m", meta=META_DU)
    wb = openpyxl.load_workbook(out)
    wb["Cover"]["C3"] = "<Project Name>"                 # tai hien dung loi goc
    wb.save(out)

    # Act
    wb2 = openpyxl.load_workbook(out)
    dat, ct = V.check_cover_toc_da_dien(V._sheet_tc(wb2), 1, None)

    # Assert
    assert not dat
    assert "Cover!C3" in ct


def test_verify_khong_soi_3_sheet_danh_cho_tester_dien(tmp_path):
    """Test report / Test data / Evidences giu placeholder la CO CHU DINH, cam fail."""
    # Arrange
    src = _csv(tmp_path, [_tc(1)])
    out = tmp_path / "o.xlsx"
    ghi(src, out, TEMPLATE, sheet_name="m", meta=META_DU)

    # Act
    wb = openpyxl.load_workbook(out)
    dat, ct = V.check_cover_toc_da_dien(V._sheet_tc(wb), 1, None)

    # Assert — 3 sheet do van con `<TC_001>`, `<Column 1>`... nhung phai PASS
    assert [wb["Test data"]["A3"].value, wb["Evidences"]["B3"].value] != [None, None]
    assert dat, ct


def test_ghi_header_sheet_test_case_dien_du_b2_b6(tmp_path):
    """B2..B6 template la `<Function Name>`..`<Link ticket>` — phai bi ghi de."""
    # Arrange
    src = _csv(tmp_path, [_tc(1)])
    out = tmp_path / "o.xlsx"

    # Act
    ghi(src, out, TEMPLATE, sheet_name="buyer_navigator_menu", meta=META_DU)

    # Assert
    ws = openpyxl.load_workbook(out)["buyer_navigator_menu"]
    assert ws["B2"].value == "buyer_navigator_menu"      # Function Name
    assert ws["B3"].value == "マイチケット"                 # Screen Name giu nguyen van
    assert ws["B4"].value == "HuongDT"
    assert ws["B6"].value == META_DU["reference"]
    sot = [o for o in ("B2", "B3", "B4", "B5", "B6")
           if V.PLACEHOLDER.search(str(ws[o].value))]
    assert sot == []


def test_ghi_header_thieu_screen_name_thi_to_do_cung_luot_voi_test_case(tmp_path):
    """Header duoc ghi TRUOC to_do_red(ws) cua writer — sai thu tu thi khong to duoc."""
    # Arrange
    from xlsx_row_ops import DO, TOKEN_CHUA_CHOT
    meta = {k: v for k, v in META_DU.items() if k != "screen_name"}
    src = _csv(tmp_path, [_tc(1)])
    out = tmp_path / "o.xlsx"

    # Act
    ghi(src, out, TEMPLATE, sheet_name="m", meta=meta)

    # Assert
    ws = openpyxl.load_workbook(out)["m"]
    assert ws["B3"].value == TOKEN_CHUA_CHOT
    assert ws["B3"].font.color.rgb == DO


def test_verify_fail_khi_header_sheet_test_case_con_placeholder(tmp_path):
    """Regression: check 12 truoc day chi soi Cover/ToC, bo lot B2..B6."""
    # Arrange
    src = _csv(tmp_path, [_tc(1)])
    out = tmp_path / "o.xlsx"
    ghi(src, out, TEMPLATE, sheet_name="m", meta=META_DU)
    wb = openpyxl.load_workbook(out)
    wb["m"]["B2"] = "<Function Name>"                    # tai hien dung loi goc
    wb.save(out)

    # Act
    wb2 = openpyxl.load_workbook(out)
    dat, ct = V.check_cover_toc_da_dien(V._sheet_tc(wb2), 1, None)

    # Assert
    assert not dat
    assert "m!B2" in ct
