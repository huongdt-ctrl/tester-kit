"""Kiem tra file test case da ghi xong. Exit 0 = dat het, 1 = co muc hong.

Dung: python3 verify_testcase_xlsx.py <file.xlsx> --template <tpl>
"""
import argparse
import re
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from xlsx_row_ops import _needed_lines, LINE_PT, DO, TOKEN_CHUA_CHOT  # noqa: E402

HEADER_ROW, DATA_FROM = 10, 12
HEADERS = {"A": "Classification", "D": "Test subject", "E": "Priority", "F": "Test case ID",
           "G": "Pre-condition", "H": "Steps to reproduce", "I": "Test data",
           "J": "Expected result"}
EXEC_COLS = [chr(ord("K") + i) for i in range(10)]
BAT_BUOC = ["D", "E", "F", "H", "J"]          # Pre-condition / Test data duoc phep rong
CAM = {"N/A", "TBD", "-", "—"}


def _sheet_tc(wb):
    """Sheet test case = sheet duy nhat co header 'Test case ID' o F10."""
    for n in wb.sheetnames:
        if wb[n][f"F{HEADER_ROW}"].value == "Test case ID":
            return wb[n]
    raise SystemExit("Khong tim thay sheet test case (F10 != 'Test case ID')")


def _so_tc(ws):
    n = 0
    while ws[f"F{DATA_FROM + n}"].value:
        n += 1
    return n


def check_header(ws, n, tpl):
    """1. Header r10 giu nguyen, khong doi ten/thu tu cot."""
    sai = [f"{c}{HEADER_ROW}={ws[f'{c}{HEADER_ROW}'].value!r}"
           for c, v in HEADERS.items() if ws[f"{c}{HEADER_ROW}"].value != v]
    return not sai, str(sai)


def check_co_testcase(ws, n, tpl):
    """2. Co it nhat 1 test case."""
    return n > 0, f"so test case = {n}"


def check_tcid_duy_nhat(ws, n, tpl):
    """3. Test case ID khong trung, khong trong."""
    ids = [ws[f"F{DATA_FROM + i}"].value for i in range(n)]
    trung = {x for x in ids if ids.count(x) > 1}
    return not trung, f"ID trung: {sorted(trung)[:5]}"


def check_cot_bat_buoc(ws, n, tpl):
    """4. Moi test case co du Test subject / Priority / ID / Steps / Expected result."""
    thieu = [f"{c}{DATA_FROM + i}" for i in range(n) for c in BAT_BUOC
             if not str(ws[f"{c}{DATA_FROM + i}"].value or "").strip()]
    return not thieu, f"o trong: {thieu[:8]}"


def check_khong_ket_qua_gia(ws, n, tpl):
    """5. Cot ket qua thi hanh K..T phai TRONG (template co san 'Passed'/'HuongDT')."""
    sot = [f"{c}{DATA_FROM + i}={ws[f'{c}{DATA_FROM + i}'].value!r}"
           for i in range(n) for c in EXEC_COLS if ws[f"{c}{DATA_FROM + i}"].value is not None]
    return not sot, f"ket qua thi hanh con sot: {sot[:5]}"


def _so_dau_dong(dong):
    """'2.1. abc' -> 2 ; '3. abc' -> 3 ; dong khong danh so -> None."""
    m = re.match(r"\s*(\d+)(?:\.\d+)*\.", dong)
    return int(m.group(1)) if m else None


def check_steps_khop_expected(ws, n, tpl):
    """6. Moi buoc trong Steps deu co Expected result tuong ung.

    KHONG so sanh so DONG: 1 buoc duoc phep co nhieu diem verify, viet phan cap
    2.1 / 2.2 / 2.3 duoi buoc 2. Chi doi hoi TAP so hieu buoc phai trung nhau.
    """
    lech = []
    for i in range(n):
        r = DATA_FROM + i
        buoc = {x for x in map(_so_dau_dong, str(ws[f"H{r}"].value or "").splitlines()) if x}
        mong = {x for x in map(_so_dau_dong, str(ws[f"J{r}"].value or "").splitlines()) if x}
        if not buoc or not mong:
            continue
        thieu, thua = sorted(buoc - mong), sorted(mong - buoc)
        if thieu or thua:
            lech.append(f"r{r}(buoc thieu expected: {thieu}; expected thua: {thua})")
    return not lech, f"{len(lech)} dong lech: {lech[:5]}"


def check_khong_cat_chu(ws, n, tpl):
    """7. Khong dong test case nao bi cat chu."""
    cat = [DATA_FROM + i for i in range(n)
           if (ws.row_dimensions[DATA_FROM + i].height or 12.75)
           < max([_needed_lines(ws, c) for c in ws[DATA_FROM + i] if c.value is not None] or [1])
           * LINE_PT * 0.85]
    return not cat, f"{len(cat)} dong bi cat: {cat[:8]}"


def check_token_cam(ws, n, tpl):
    """8. Khong dung N/A / TBD / - lam gia tri thieu."""
    hits = [f"{c}{DATA_FROM + i}" for i in range(n) for c in HEADERS
            if str(ws[f"{c}{DATA_FROM + i}"].value or "").strip() in CAM]
    return not hits, f"token cam: {hits[:8]}"


def check_expected_khong_chua_chot(ws, n, tpl):
    """10. Expected result KHONG duoc chua 'Can xac nhan' (SKILL.md §10.3).

    Test case co ket qua mong doi chua chot thi khong chay duoc -> cam ban giao.
    Khac gen-test-plan: o do 'Can xac nhan' la hop le va chi can to do.
    """
    xau = [f"J{DATA_FROM + i}" for i in range(n)
           if TOKEN_CHUA_CHOT in str(ws[f"J{DATA_FROM + i}"].value or "")]
    return not xau, f"{len(xau)} case co Expected result chua chot: {xau[:8]}"


def check_chua_chot_da_to_do(ws, n, tpl):
    """11. O nao lot 'Can xac nhan' thi phai to chu do de nguoi review thay.

    Xet MOI lan xuat hien (khac /gen-test-plan): noi dung test case khong co van
    mo ta nhac lai cum tu, nen token o dau cung la diem can nguoi xu ly.
    """
    xau = []
    for i in range(n):
        for c in HEADERS:
            cell = ws[f"{c}{DATA_FROM + i}"]
            if TOKEN_CHUA_CHOT not in str(cell.value or ""):
                continue
            mau = getattr(cell.font.color, "rgb", None) if cell.font and cell.font.color else None
            if mau != DO:
                xau.append(f"{c}{DATA_FROM + i}")
    return not xau, f"{len(xau)} o chua to do: {xau[:8]}"


def check_khop_template(ws, n, tpl):
    """9. So sheet va so cot khong doi so voi template."""
    if tpl is None:
        return True, "(bo qua)"
    loi = []
    if len(tpl.sheetnames) != len(ws.parent.sheetnames):
        loi.append(f"so sheet {len(tpl.sheetnames)} -> {len(ws.parent.sheetnames)}")
    if tpl["Function {Name}"].max_column != ws.max_column:
        loi.append(f"so cot {tpl['Function {Name}'].max_column} -> {ws.max_column}")
    return not loi, str(loi)


CHECKS = [check_header, check_co_testcase, check_tcid_duy_nhat, check_cot_bat_buoc,
          check_khong_ket_qua_gia, check_steps_khop_expected, check_khong_cat_chu,
          check_token_cam, check_expected_khong_chua_chot,
          check_chua_chot_da_to_do, check_khop_template]


def run(path, template=None):
    wb = openpyxl.load_workbook(path)
    ws = _sheet_tc(wb)
    n = _so_tc(ws)
    tpl = openpyxl.load_workbook(template) if template else None
    fail = 0
    for fn in CHECKS:
        dat, ct = fn(ws, n, tpl)
        print(f"  {'PASS' if dat else 'FAIL'}  {(fn.__doc__ or '').strip().splitlines()[0]}")
        if not dat:
            print(f"        {ct}")
            fail += 1
    print(f"\n  {len(CHECKS) - fail} pass, {fail} fail  ({n} test case, sheet {ws.title!r})")
    return fail


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx"); ap.add_argument("--template")
    a = ap.parse_args()
    sys.exit(1 if run(a.xlsx, a.template) else 0)
