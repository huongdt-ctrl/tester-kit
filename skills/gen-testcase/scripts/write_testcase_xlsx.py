"""Ghi test case tu CSV vao ban copy cua template_testcase.xlsx.

Dung: python3 write_testcase_xlsx.py <csv> <out.xlsx> --template <tpl> [--sheet-name <ten>]

Vi sao co script nay: truoc day skill chi tao Google Sheet, gap loi vo format va
phai co Drive auth, nen test case ket lai o CSV. Script ghi thang ra .xlsx local.
"""
import argparse
import csv
import re
import shutil
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment
from openpyxl.utils.cell import range_boundaries

sys.path.insert(0, str(Path(__file__).resolve().parent))
from xlsx_row_ops import autofit_rows, insert_rows_keep_merges, to_do_red  # noqa: E402

SHEET_MAU = "Function {Name}"
HEADER_ROW = 10                 # r10 = ten cot, r11 = nhan I/II/III cho A/B/C
DATA_FROM = 12
DATA_TO = 20                    # template chi ke san 9 dong (12-20)
COLS = ["A", "B", "C", "D", "E", "F", "G", "H", "I", "J"]
CSV_COLS = ["Classification I", "Classification II", "Classification III", "Test subject",
            "Priority", "Test case ID", "Pre-condition", "Steps to reproduce",
            "Test data", "Expected result"]
WRAP_COLS = ["D", "G", "H", "I", "J"]   # template chi bat wrap cho I/J -> thieu D/G/H
EXEC_COLS = [chr(ord("K") + i) for i in range(10)]   # K..T: ket qua thi hanh Round 1/2


def _ten_sheet_hop_le(ten):
    """Excel: toi da 31 ky tu, cam [ ] : * ? / \\ ."""
    ten = re.sub(r"[\[\]:*?/\\]", "_", str(ten)).strip() or "Testcases"
    return ten[:31]


def doc_csv(path):
    with open(path, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    if not rows:
        raise SystemExit(f"CSV rong: {path}")
    thieu = [c for c in CSV_COLS if c not in rows[0]]
    if thieu:
        raise SystemExit(f"CSV thieu cot: {thieu}\nCot dang co: {list(rows[0])}")
    return rows


def xoa_du_lieu_mau(ws, tu, den):
    """Template co san TC_001.. va CA KET QUA THI HANH 'Passed'/'HuongDT'/2023-07-07.

    Khong xoa thi ban giao test case kem ket qua pass gia -- loi nghiem trong.
    """
    for m in [str(m) for m in ws.merged_cells.ranges]:
        c1, r1, c2, r2 = range_boundaries(m)
        if r1 >= tu and r2 <= den and r2 > r1:
            ws.unmerge_cells(m)                    # bo merge doc cua A/B/C mau
    for r in range(tu, den + 1):
        for col in COLS + EXEC_COLS:
            ws[f"{col}{r}"].value = None


def gom_nhom_cot_phan_loai(ws, tu, den):
    """Merge doc A/B/C cho cac dong lien tiep cung gia tri -- dung cach template trinh bay."""
    da_merge = 0
    for col in ("A", "B", "C"):
        r = tu
        while r <= den:
            v = ws[f"{col}{r}"].value
            if v is None:
                r += 1
                continue
            cuoi = r
            while cuoi + 1 <= den and ws[f"{col}{cuoi + 1}"].value == v:
                cuoi += 1
            if cuoi > r:
                for x in range(r + 1, cuoi + 1):
                    ws[f"{col}{x}"].value = None   # giu lai o dau, xoa o lap
                ws.merge_cells(start_row=r, start_column=ws[f"{col}{r}"].column,
                               end_row=cuoi, end_column=ws[f"{col}{r}"].column)
                ws[f"{col}{r}"].alignment = Alignment(
                    horizontal="left", vertical="top", wrap_text=True)
                da_merge += 1
            r = cuoi + 1
    return da_merge


def ghi(csv_path, out_path, template, sheet_name=None):
    rows = doc_csv(csv_path)
    shutil.copyfile(template, out_path)
    wb = openpyxl.load_workbook(out_path)
    ws = wb[SHEET_MAU]

    thuc_te = ws[f"D{HEADER_ROW}"].value
    if thuc_te != "Test subject":                  # §11: verify header that truoc khi ghi
        raise SystemExit(f"HEADER LECH: D{HEADER_ROW} = {thuc_te!r}, mong doi 'Test subject'")

    ws.title = _ten_sheet_hop_le(sheet_name or Path(csv_path).parent.name)
    xoa_du_lieu_mau(ws, DATA_FROM, DATA_TO)

    can = len(rows)
    co = DATA_TO - DATA_FROM + 1
    if can > co:                                   # chen dong, copy format tu dong lien tren
        insert_rows_keep_merges(ws, at=DATA_TO, n=can - co, clone_from=DATA_TO - 1)
    den = DATA_FROM + can - 1

    for i, rec in enumerate(rows):
        r = DATA_FROM + i
        for col, key in zip(COLS, CSV_COLS):
            gt = (rec.get(key) or "").strip()
            if gt:
                ws[f"{col}{r}"] = gt
            if col in WRAP_COLS:
                ws[f"{col}{r}"].alignment = Alignment(
                    horizontal="left", vertical="top", wrap_text=True)

    # THU TU QUAN TRONG: autofit TRUOC, gom nhom SAU.
    # autofit_rows bo qua dong nam trong merge doc; merge A/B/C truoc thi moi dong
    # data deu nam trong merge => khong dong nao duoc chinh chieu cao.
    autofit_rows(ws, range(DATA_FROM, den + 1))
    nhom = gom_nhom_cot_phan_loai(ws, DATA_FROM, den)
    # Luoi an toan: 'Can xac nhan' bi CAM trong noi dung test case (SKILL.md §10.3).
    # Neu lot vao thi to do de nguoi review thay ngay; verify se chan viec ban giao.
    to_do_red(ws, bat_ky=True)   # trong test case, moi lan xuat hien deu la loi
    wb.save(out_path)
    return len(rows), nhom, ws.title


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("csv"); ap.add_argument("out")
    ap.add_argument("--template", required=True); ap.add_argument("--sheet-name")
    a = ap.parse_args()
    n, nhom, ten = ghi(a.csv, a.out, a.template, a.sheet_name)
    print(f"OK — {n} test case -> sheet {ten!r} · {nhom} nhom phan loai da merge\n   {a.out}")
