# -*- coding: utf-8 -*-
"""Thao tac dong tren xlsx giu nguyen merge / style / row height.

openpyxl 3.1.5 `insert_rows` KHONG dich merged range, data validation va row height
=> phai tu xu ly, neu khong bang bi vo merge.
"""
import math
from copy import copy
from openpyxl.utils.cell import range_boundaries, get_column_letter

DEFAULT_COL_WIDTH = 8.43
LINE_PT = 14.5          # chieu cao 1 dong text font ~11pt
MAX_ROW_PT = 409        # gioi han cua Excel


def _row_merges(ws, row):
    """Cac merge nam tron trong 1 dong (dung de nhan ban mau dong)."""
    out = []
    for m in ws.merged_cells.ranges:
        c1, r1, c2, r2 = range_boundaries(str(m))
        if r1 == row and r2 == row:
            out.append((c1, c2))
    return out


def insert_rows_keep_merges(ws, at, n, clone_from):
    """Chen n dong tai vi tri `at`, giu merge cu, roi nhan ban style+merge cua dong `clone_from`.

    `clone_from` tinh theo he toa do TRUOC khi chen.
    """
    if n <= 0:
        return
    saved = [str(m) for m in ws.merged_cells.ranges]
    for m in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(m))
    heights = {r: d.height for r, d in ws.row_dimensions.items() if d.height is not None}

    ws.insert_rows(at, n)

    for m in saved:                                  # dich merge cu
        c1, r1, c2, r2 = range_boundaries(m)
        if r1 >= at:
            r1 += n; r2 += n
        elif r2 >= at:
            r2 += n                                  # vung bac qua diem chen -> keo dai
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)

    for r in sorted(heights, reverse=True):          # dich row height
        h = heights[r]
        ws.row_dimensions[r + n if r >= at else r].height = h

    src = clone_from + n if clone_from >= at else clone_from
    pattern = _row_merges(ws, src)
    for i in range(n):                               # nhan ban mau dong
        dst = at + i
        for c in range(1, ws.max_column + 1):
            s, d = ws.cell(row=src, column=c), ws.cell(row=dst, column=c)
            d._style = copy(s._style)
        for c1, c2 in pattern:
            if c1 != c2:
                ws.merge_cells(start_row=dst, start_column=c1, end_row=dst, end_column=c2)
        ws.row_dimensions[dst].height = ws.row_dimensions[src].height


def normalize_row_merges(ws, first_row, n):
    """Nhan ban mau merge cua dong data dau tien ra n dong.

    Template TEM-ST02-01 chi merge ~5 dong dau moi bang; cac dong sau chi co border
    ma khong merge => Excel hien duong ke doc BEN TRONG o, nhin lech han cac dong tren.
    """
    pattern = _row_merges(ws, first_row)
    if not pattern:
        return 0
    added = 0
    for i in range(n):
        r = first_row + i
        existing = _row_merges(ws, r)
        for c1, c2 in pattern:
            if c1 == c2:
                continue
            if any(a <= c2 and c1 <= b for a, b in existing):   # da co merge chong lan
                continue
            ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
            added += 1
    return added


def _visual_len(s):
    """Ky tu CJK chiem 2 o chieu rong."""
    return sum(2 if ord(ch) > 0x2E7F else 1 for ch in s)


def _span_width(ws, cell):
    """Chieu rong (don vi ky tu) cua o, cong ca vung merge ngang."""
    c1 = c2 = cell.column
    for m in ws.merged_cells.ranges:
        a1, r1, a2, r2 = range_boundaries(str(m))
        if r1 <= cell.row <= r2 and a1 <= cell.column <= a2:
            c1, c2 = a1, a2
            break
    total = 0.0
    for c in range(c1, c2 + 1):
        d = ws.column_dimensions.get(get_column_letter(c))
        total += (d.width if d and d.width else DEFAULT_COL_WIDTH)
    return max(total, 4.0)


def _needed_lines(ws, cell):
    v = cell.value
    if not isinstance(v, str) or not v.strip():
        return 1
    if not (cell.alignment and cell.alignment.wrap_text):
        return v.count("\n") + 1              # khong wrap => chi xuong dong o ky tu \n
    per = _span_width(ws, cell)
    return sum(max(1, math.ceil(_visual_len(seg) / per)) for seg in v.split("\n"))


def autofit_rows(ws, rows):
    """Dat chieu cao dong theo so dong text thuc te sau khi wrap.

    O nam trong merge DOC: tinh tong chieu cao ca block, thieu bao nhieu thi chia deu.
    """
    vblocks = {}
    for m in ws.merged_cells.ranges:
        c1, r1, c2, r2 = range_boundaries(str(m))
        if r2 > r1:
            for r in range(r1, r2 + 1):
                vblocks[r] = (r1, r2)

    changed = 0
    need_single = {}
    need_block = {}
    for r in rows:
        for cell in ws[r]:
            if cell.value is None:
                continue
            n = _needed_lines(ws, cell)
            if n < 2:
                continue
            blk = vblocks.get(cell.row)
            if blk and blk[0] <= cell.row <= blk[1] and blk[1] > blk[0]:
                need_block[blk] = max(need_block.get(blk, 1), n)
            else:
                need_single[r] = max(need_single.get(r, 1), n)

    for r, n in need_single.items():
        h = min(n * LINE_PT, MAX_ROW_PT)
        if (ws.row_dimensions[r].height or 0) < h:
            ws.row_dimensions[r].height = h
            changed += 1

    for (r1, r2), n in need_block.items():           # chia deu phan thieu cho ca block
        have = sum(ws.row_dimensions[r].height or 12.75 for r in range(r1, r2 + 1))
        want = min(n * LINE_PT, MAX_ROW_PT * (r2 - r1 + 1))
        if have >= want:
            continue
        extra = (want - have) / (r2 - r1 + 1)
        for r in range(r1, r2 + 1):
            ws.row_dimensions[r].height = min((ws.row_dimensions[r].height or 12.75) + extra, MAX_ROW_PT)
            changed += 1
    return changed


def ensure_narrative_merge(ws, addr, end_col):
    """O narrative dai nhung KHONG merge => wrap se ep vao be rong 1 cot (~8 ky tu).

    Template TEM-ST02-01 khong nhat quan: B10:W11, B3:X17 co merge, B6 thi khong.
    Ham nay merge bo sung cho dong nhat, KHONG dung vao o da co merge san.
    """
    cell = ws[addr]
    for m in ws.merged_cells.ranges:
        if cell.coordinate in m:
            return False
    from openpyxl.utils.cell import column_index_from_string
    c2 = column_index_from_string(end_col)
    if c2 <= cell.column:
        return False
    ws.merge_cells(start_row=cell.row, start_column=cell.column, end_row=cell.row, end_column=c2)
    return True


def find_header_row(ws, col, text, lo, hi):
    """Do lai dong header sau khi chen dong (§11 verify header truoc khi ghi)."""
    for r in range(lo, hi + 1):
        v = ws[f"{col}{r}"].value
        if v and text in str(v):
            return r
    raise SystemExit(f"KHONG THAY header {text!r} o cot {col} trong {lo}-{hi} cua {ws.title}")

def anchor(ws, addr):
    """O nam trong vung merge -> tra ve o goc tren trai (o duy nhat ghi duoc).

    Khong dung ham nay thi `ws[addr].value = x` nem
    `AttributeError: MergedCell object attribute is read-only`.
    Sheet 07 co A5:A19 va B5:B7 / B8:B10 / B11:B15 / B16:B18 (xem cell map B6).
    """
    from openpyxl.cell.cell import MergedCell
    cell = ws[addr]
    if isinstance(cell, MergedCell):
        for rng in ws.merged_cells.ranges:
            if cell.coordinate in rng:
                return ws.cell(row=rng.min_row, column=rng.min_col)
    return cell


def sync_number_format(ws, col, rows, ref_row):
    """Dong bo number_format cua `col` tai `rows` theo dong `ref_row` cung cot.

    Template co loi: 07!G5/H5 la '0.00' nhung G6:H19 la 'yyyy/mm/dd' (cell map B4).
    Ghi so gio vao do thi Excel hien thanh ngay/gio ma khong bao loi.
    """
    ref = ws[f"{col}{ref_row}"].number_format
    changed = 0
    for r in rows:
        cell = anchor(ws, f"{col}{r}")
        if cell.number_format != ref:
            cell.number_format = ref
            changed += 1
    return changed
