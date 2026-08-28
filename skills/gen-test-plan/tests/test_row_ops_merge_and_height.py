"""Test cho scripts/xlsx_row_ops.py — thao tac dong tren xlsx giu nguyen merge/style.

Vi sao can: openpyxl `insert_rows` KHONG dich merged range / row height / data validation.
Mat merge => bang vo trinh bay ma khong co loi nao bao ra. 4 ham duoi day la thu chan viec do.
"""
import sys
from pathlib import Path

import openpyxl
import pytest
from openpyxl.utils.cell import range_boundaries

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

from xlsx_row_ops import (  # noqa: E402
    autofit_rows,
    ensure_narrative_merge,
    find_header_row,
    insert_rows_keep_merges,
    normalize_row_merges,
)


def _sheet_with_table(n_rows=3):
    """Bang mo phong template: header dong 1, data tu dong 2, moi dong merge C:E va F:H."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["C1"], ws["F1"] = "Ten muc", "Mo ta"
    for r in range(2, 2 + n_rows):
        ws[f"C{r}"] = f"muc {r}"
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
        ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=8)
    return ws


def _row_merges(ws, row):
    return sorted(
        (c1, c2)
        for m in ws.merged_cells.ranges
        for c1, r1, c2, r2 in [range_boundaries(str(m))]
        if r1 == row == r2
    )


def test_insert_rows_giu_nguyen_merge_cua_dong_phia_duoi():
    # Arrange — 3 dong data, dong 4 la dong cuoi
    ws = _sheet_with_table(3)
    truoc = _row_merges(ws, 4)

    # Act — chen 2 dong vao giua
    insert_rows_keep_merges(ws, at=3, n=2, clone_from=2)

    # Assert — dong 4 cu da xuong dong 6, merge di theo
    assert _row_merges(ws, 6) == truoc
    assert ws["C6"].value == "muc 4"


def test_insert_rows_nhan_ban_mau_merge_cho_dong_moi():
    # Arrange
    ws = _sheet_with_table(3)
    mau = _row_merges(ws, 2)

    # Act
    insert_rows_keep_merges(ws, at=3, n=2, clone_from=2)

    # Assert — 2 dong moi (3, 4) co dung mau merge cua dong 2
    assert _row_merges(ws, 3) == mau
    assert _row_merges(ws, 4) == mau


def test_insert_rows_giu_nguyen_chieu_cao_dong_bi_day_xuong():
    # Arrange
    ws = _sheet_with_table(3)
    ws.row_dimensions[4].height = 90.0

    # Act
    insert_rows_keep_merges(ws, at=3, n=2, clone_from=2)

    # Assert
    assert ws.row_dimensions[6].height == 90.0


def test_normalize_row_merges_bu_merge_cho_dong_thieu():
    """Template that chi merge ~5 dong dau moi bang; dong sau chi co border."""
    # Arrange — dong 2 co merge, dong 3 va 4 khong
    ws = _sheet_with_table(1)
    ws["C3"], ws["C4"] = "muc 3", "muc 4"
    assert _row_merges(ws, 3) == []

    # Act
    added = normalize_row_merges(ws, first_row=2, n=3)

    # Assert
    assert added == 4  # 2 merge x 2 dong con thieu
    assert _row_merges(ws, 3) == _row_merges(ws, 2) == _row_merges(ws, 4)


def test_normalize_row_merges_khong_dung_vao_dong_da_co_merge():
    # Arrange
    ws = _sheet_with_table(3)

    # Act — goi lai lan 2 tren bang da day du merge
    added = normalize_row_merges(ws, first_row=2, n=3)

    # Assert — khong them gi, khong ném loi trung merge
    assert added == 0


def test_ensure_narrative_merge_merge_o_chua_merge():
    # Arrange
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["B2"] = "doan van dai"

    # Act
    ket_qua = ensure_narrative_merge(ws, "B2", "H")

    # Assert
    assert ket_qua is True
    assert _row_merges(ws, 2) == [(2, 8)]


def test_ensure_narrative_merge_bo_qua_o_da_co_merge_san():
    """Template co san B10:W11 — cham vao se hong layout."""
    # Arrange
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.merge_cells("B2:D2")

    # Act
    ket_qua = ensure_narrative_merge(ws, "B2", "H")

    # Assert — giu nguyen merge cu
    assert ket_qua is False
    assert _row_merges(ws, 2) == [(2, 4)]


def test_autofit_rows_nang_chieu_cao_theo_so_dong_text():
    # Arrange — o rong 10 ky tu, noi dung 3 doan tach bang \n
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.column_dimensions["B"].width = 10
    ws["B2"] = "dong mot\ndong hai\ndong ba"
    ws["B2"].alignment = openpyxl.styles.Alignment(wrap_text=True)

    # Act
    autofit_rows(ws, [2])

    # Assert — it nhat 3 dong text
    assert ws.row_dimensions[2].height >= 3 * 14.0


def test_autofit_rows_bo_qua_dong_khong_bat_wrap():
    # Arrange
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["B2"] = "x" * 500          # dai nhung khong wrap => chi 1 dong hien thi

    # Act
    autofit_rows(ws, [2])

    # Assert
    assert ws.row_dimensions[2].height is None


def test_autofit_rows_chia_deu_chieu_cao_cho_merge_doc():
    # Arrange — o B2 trai dai 2 dong
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.column_dimensions["B"].width = 10
    ws.merge_cells("B2:B3")
    ws["B2"] = "\n".join(f"dong {i}" for i in range(6))
    ws["B2"].alignment = openpyxl.styles.Alignment(wrap_text=True)

    # Act
    autofit_rows(ws, [2, 3])

    # Assert — tong chieu cao ca block du cho 6 dong, khong don het vao 1 dong
    tong = ws.row_dimensions[2].height + ws.row_dimensions[3].height
    assert tong >= 6 * 14.0
    assert ws.row_dimensions[3].height > 12.75


def test_find_header_row_tra_ve_dong_khop_chuoi():
    # Arrange
    ws = _sheet_with_table(2)

    # Act
    r = find_header_row(ws, "C", "Ten muc", 1, 10)

    # Assert
    assert r == 1


def test_find_header_row_dung_han_khi_khong_thay_header():
    """Cell map lech => phai dung, cam ghi mo (§11 SKILL.md)."""
    # Arrange
    ws = _sheet_with_table(2)

    # Act + Assert
    with pytest.raises(SystemExit):
        find_header_row(ws, "C", "Header khong ton tai", 1, 10)
