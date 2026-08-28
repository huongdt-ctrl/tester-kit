"""Test cho scripts/verify_testplan.py va 2 helper anchor / sync_number_format.

Vi sao can: verify_testplan.py la thu CHAN — neu no luon PASS thi vo dung.
Cac test duoi day chung minh no fail dung luc.
"""
import sys
from pathlib import Path

import openpyxl
import pytest

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
TEMPLATE = Path(__file__).resolve().parents[1] / "templates" / "template_testplan.xlsx"
sys.path.insert(0, str(SCRIPTS))

from xlsx_row_ops import anchor, sync_number_format  # noqa: E402
import verify_testplan as V  # noqa: E402


# ---------- anchor ----------

def test_anchor_tra_ve_o_goc_khi_o_nam_trong_merge():
    # Arrange
    wb = openpyxl.Workbook(); ws = wb.active
    ws.merge_cells("A5:A19")

    # Act
    c = anchor(ws, "A12")

    # Assert
    assert c.coordinate == "A5"
    c.value = "Release 1"          # phai ghi duoc, khong nem read-only
    assert ws["A5"].value == "Release 1"


def test_anchor_tra_ve_chinh_o_khi_khong_merge():
    # Arrange
    wb = openpyxl.Workbook(); ws = wb.active

    # Act + Assert
    assert anchor(ws, "C3").coordinate == "C3"


def test_ghi_thang_vao_merged_cell_van_nem_loi_neu_khong_dung_anchor():
    """Chung minh ly do anchor() ton tai."""
    # Arrange
    wb = openpyxl.Workbook(); ws = wb.active
    ws.merge_cells("A5:A19")

    # Act + Assert
    with pytest.raises(AttributeError):
        ws["A12"].value = "x"


# ---------- sync_number_format ----------

def test_sync_number_format_sua_o_bi_format_ngay_thanh_so():
    """Template co loi: G5 la '0.00' nhung G6:G19 la 'yyyy/mm/dd' (cell map B4)."""
    # Arrange
    wb = openpyxl.Workbook(); ws = wb.active
    ws["G5"].number_format = "0.00"
    for r in range(6, 10):
        ws[f"G{r}"].number_format = "yyyy/mm/dd"

    # Act
    changed = sync_number_format(ws, "G", range(6, 10), ref_row=5)

    # Assert
    assert changed == 4
    assert {ws[f"G{r}"].number_format for r in range(5, 10)} == {"0.00"}


def test_sync_number_format_khong_dong_gi_khi_da_dung():
    # Arrange
    wb = openpyxl.Workbook(); ws = wb.active
    for r in range(5, 10):
        ws[f"G{r}"].number_format = "0.00"

    # Act + Assert
    assert sync_number_format(ws, "G", range(6, 10), ref_row=5) == 0


# ---------- verify_testplan ----------

def test_verify_bat_loi_tren_template_chua_ghi_gi():
    """Template trang phai FAIL — neu PASS thi script la con dau cao su."""
    # Arrange
    wb = openpyxl.load_workbook(TEMPLATE)

    # Act
    ket_qua = [fn(wb, None)[0] for fn in V.CHECKS]

    # Assert
    assert not all(ket_qua), "verify PASS het tren template trang => vo dung"


def test_check_rui_ro_fail_khi_thieu_bien_phap():
    # Arrange — co mo ta rui ro nhung bo trong cot R
    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb["09_Risk management"]
    for r in range(11, 27):
        ws[f"P{r}"], ws[f"Q{r}"] = "Cao", "Mo"

    # Act
    dat, chi_tiet = V.check_rui_ro(wb, None)

    # Assert
    assert dat is False and "R11" in chi_tiet


def test_check_nguong_criteria_fail_khi_tieu_chi_khong_co_so():
    # Arrange
    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb["06_Criteria"]
    ws["C8"], ws["F8"] = "Chat luong", "Test day du, khach hang hai long"

    # Act
    dat, _ = V.check_nguong_criteria(wb, None)

    # Assert
    assert dat is False


def test_check_nguong_criteria_pass_khi_co_ky_hieu_so():
    # Arrange
    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb["06_Criteria"]
    ws["C8"], ws["F8"] = "Ty le pass", "Ty le test case pass >= 95%."

    # Act
    dat, _ = V.check_nguong_criteria(wb, None)

    # Assert
    assert dat is True


def test_check_token_cam_bo_qua_dau_gach_von_co_trong_template():
    """01_Introduction!B5 = '-' la gach dau dong cua template, khong phai gia tri thieu."""
    # Arrange
    wb = openpyxl.load_workbook(TEMPLATE)
    tpl = openpyxl.load_workbook(TEMPLATE)

    # Act
    dat, _ = V.check_token_cam(wb, tpl)

    # Assert
    assert dat is True


def test_check_token_cam_bat_khi_buoc_ghi_dat_N_A():
    # Arrange
    wb = openpyxl.load_workbook(TEMPLATE)
    tpl = openpyxl.load_workbook(TEMPLATE)
    wb["02_Scope test"]["C30"] = "N/A"

    # Act
    dat, chi_tiet = V.check_token_cam(wb, tpl)

    # Assert
    assert dat is False and "C30" in chi_tiet
