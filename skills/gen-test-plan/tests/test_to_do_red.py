"""Test cho to_do_red / la_diem_chua_chot — to do chu o diem chua chot."""
import sys
from pathlib import Path

import openpyxl
import pytest
from openpyxl.styles import Font

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))

from xlsx_row_ops import DO, la_diem_chua_chot, to_do_red  # noqa: E402
import verify_testplan as V  # noqa: E402


@pytest.mark.parametrize("gia_tri,mong_doi", [
    ("Cần xác nhận", True),
    ("Cần xác nhận — chưa có quyền truy cập", True),
    ("Cần xác nhận\nManifest mục run.reviewer đang trống", True),
    ("Người có thẩm quyền: Cần xác nhận", True),
    ("Mức tương đồng với production: Cần xác nhận.", True),
    ("Mục 5.1, 5.2, 5.3 phải ghi Cần xác nhận cho số lượng", False),
    ("mục 5 hiện còn nhiều điểm Cần xác nhận.", False),
    ("Full test case", False),
    ("", False),
    (None, False),
    (123, False),
])
def test_la_diem_chua_chot_phan_loai_dung(gia_tri, mong_doi):
    # Arrange + Act + Assert
    assert la_diem_chua_chot(gia_tri) is mong_doi


def test_to_do_red_to_o_bat_dau_bang_token():
    # Arrange
    wb = openpyxl.Workbook(); ws = wb.active
    ws["B2"] = "Cần xác nhận — thiếu tài liệu môi trường"

    # Act
    n = to_do_red(ws)

    # Assert
    assert n == 1
    assert ws["B2"].font.color.rgb == DO


def test_to_do_red_khong_to_o_nhac_token_giua_cau():
    """O mo ta quy tac (vd 'phai ghi Can xac nhan cho...') khong phai diem chua chot."""
    # Arrange
    wb = openpyxl.Workbook(); ws = wb.active
    ws["B2"] = "Mục 5.1, 5.2, 5.3 phải ghi Cần xác nhận cho số lượng và version"

    # Act
    n = to_do_red(ws)

    # Assert
    assert n == 0
    assert ws["B2"].font.color is None or ws["B2"].font.color.rgb != DO


def test_to_do_red_giu_nguyen_thuoc_tinh_font_khac():
    # Arrange
    wb = openpyxl.Workbook(); ws = wb.active
    ws["B2"] = "Cần xác nhận"
    ws["B2"].font = Font(name="Arial", size=10, bold=True, italic=True)

    # Act
    to_do_red(ws)

    # Assert — chi doi mau
    f = ws["B2"].font
    assert (f.name, f.size, f.bold, f.italic) == ("Arial", 10, True, True)
    assert f.color.rgb == DO


def test_to_do_red_ghi_qua_o_goc_khi_nam_trong_merge():
    # Arrange
    wb = openpyxl.Workbook(); ws = wb.active
    ws.merge_cells("B2:E2")
    ws["B2"] = "Cần xác nhận"

    # Act
    n = to_do_red(ws)

    # Assert
    assert n == 1 and ws["B2"].font.color.rgb == DO


def test_verify_bat_o_chua_chot_chua_to_do():
    # Arrange
    wb = openpyxl.Workbook()
    wb.active["B2"] = "Cần xác nhận"

    # Act
    dat, ct = V.check_can_xac_nhan_mau_do(wb, None)

    # Assert
    assert dat is False and "B2" in ct


def test_verify_pass_sau_khi_to_do():
    # Arrange
    wb = openpyxl.Workbook(); ws = wb.active
    ws["B2"] = "Cần xác nhận"
    to_do_red(ws)

    # Act
    dat, _ = V.check_can_xac_nhan_mau_do(wb, None)

    # Assert
    assert dat is True
