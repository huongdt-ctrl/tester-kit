#!/usr/bin/env python3
"""Clone template TEM-ST03_02 -> file phan tich bug cua mot giai doan.

Thu tu BAT BUOC (doi thu tu la hong file):
  1. Tinh Layout tu so bug / so man hinh / so milestone that
  2. Chen dong (sheet_ops) -- phai lam TRUOC khi ghi bat cu o nao, vi chen dong
     lam moi toa do ben duoi thay doi
  3. Xoa du lieu mau cua template
  4. Ghi du lieu bug + nhan block dong
  5. Sinh lai cong thuc theo toa do MOI
  6. Va lai dropdown + chart
"""
import os
import shutil

import openpyxl

import catalogue
import formula_writer as fw
import sheet_ops
from template_layout import BLOCKS, SHEET_NAME, Layout


def output_filename(phase, date_from, date_to):
    """<Giai doan>_Phan tich Bug (co dau)_<YYYYMMDD-YYYYMMDD>.xlsx (quy uoc da chot)."""
    span = "%s-%s" % (_compact(date_from), _compact(date_to))
    return "%s_Phân tích Bug_%s.xlsx" % (_safe(phase or "Giai doan"), span)


def _compact(value):
    return str(value or "").replace("-", "").strip() or "00000000"


def _safe(value):
    bad = '/\\:*?"<>|'
    return "".join(ch for ch in str(value) if ch not in bad).strip()


def distinct_ordered(rows, key):
    """Nhan cho bang thong ke dong, giu thu tu xuat hien dau tien.

    Gop khong phan biet hoa/thuong: COUNTIF cua Excel **bo qua hoa/thuong**, nen
    de "G10" va "g10" thanh hai dong nhan thi moi dong dem ca hai -> TOTAL gan
    nhu dem doi. Giu cach viet gap dau tien lam nhan chinh thuc.
    """
    seen, out = set(), []
    for item in rows:
        value = (item.get(key) or "").strip()
        folded = value.lower()
        if value and folded not in seen:
            seen.add(folded)
            out.append(value)
    return out


def build(template_path, out_path, rows, phase=None, date_from=None,
          date_to=None, issues=None, overwrite=False):
    """Tao file phan tich bug. Tra ve dict bao cao de CLI in ra JSON."""
    if not os.path.exists(template_path):
        raise FileNotFoundError("Khong tim thay template: %s" % template_path)
    if os.path.exists(out_path) and not overwrite:
        raise FileExistsError(
            "File da ton tai: %s -- them --overwrite neu that su muon ghi de" % out_path)
    os.makedirs(os.path.dirname(os.path.abspath(out_path)) or ".", exist_ok=True)
    shutil.copyfile(template_path, out_path)

    wb = openpyxl.load_workbook(out_path)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError("Template thieu sheet %r (co: %s)"
                         % (SHEET_NAME, ", ".join(wb.sheetnames)))
    ws = wb[SHEET_NAME]

    labels, added = _plan_labels(wb, ws, rows)
    layout = Layout(len(rows), {key: len(val) for key, val in labels.items()})

    sheet_ops.apply_insertions(ws, layout)
    fw.clear_data_region(ws, layout)
    fw.write_data_rows(ws, layout, rows)
    fw.write_phase_header(ws, phase, date_from, date_to)

    for block in layout.blocks():
        fw.write_block(ws, layout, block, labels.get(block["key"]))

    validated = sheet_ops.remap_validations(ws, layout)
    chart_fixes = sheet_ops.remap_charts(ws, layout, SHEET_NAME)
    fw.write_issue_actions(ws, layout, issues)

    wb.save(out_path)
    return {
        "output": out_path,
        "bug_count": len(rows),
        "rows_inserted": dict(layout.extra),
        "data_range": "%s3:%s%d" % ("A", "M", layout.data_end),
        "screens": labels.get("screen", []),
        "milestones": labels.get("milestone", []),
        "validations_extended": validated,
        "chart_refs_fixed": chart_fixes,
        "catalogue_rows_added": added,
        "uncounted": reconcile(rows, labels),
    }


def _plan_labels(wb, ws, rows):
    """Nhan cua ca 6 bang thong ke, tinh o TOA DO GOC (truoc khi chen dong).

    Bang dong (man hinh / milestone) lay nhan tu du lieu that. Bang danh muc co
    dinh lay nhan dang co, roi NOI THEM nhung muc ma sheet 'Define' co ma bang
    bo sot -- neu khong, bug xep vao muc bi sot se khong duoc dem vao dau ca.
    """
    base = Layout()
    full = catalogue.read(wb)
    labels = {"screen": distinct_ordered(rows, "screen"),
              "milestone": distinct_ordered(rows, "milestone")}
    added = {}
    for block in BLOCKS:
        if block.dynamic:
            continue
        existing = fw.read_fixed_labels(ws, base.block(block.key))
        completed = catalogue.complete(existing, full.get(block.key))
        labels[block.key] = completed
        missing = [lb for lb in completed
                   if lb not in {str(v).strip() for v in existing if v}]
        if missing:
            added[block.key] = missing
    return labels, added


def reconcile(rows, fixed_labels):
    """Bug co gia tri KHONG nam trong danh muc cua bang thong ke -> se khong duoc
    dem vao dau ca. Bao ra thay vi de TOTAL lech am tham.

    Template co khoang trong that: dropdown Cause co 22 lua chon nhung bang VI
    chi liet ke 17 (thieu cac muc '... Other'). Chon nham vao 5 muc do thi bug
    bien mat khoi thong ke.
    """
    pairs = (("cause", "cause"), ("root_cause", "rootcause"),
             ("priority", "priority"), ("severity", "severity"))
    out = []
    for field, block_key in pairs:
        # So khong phan biet hoa/thuong: COUNTIF cua Excel dem duoc "major" cho
        # nhan "Major". Neu o day so case-sensitive thi bao nham la bug do
        # khong duoc dem, agent lai di "sua" mot gia tri von dang dung.
        allowed = {str(v).strip().lower() for v in (fixed_labels.get(block_key) or []) if v}
        if not allowed:
            continue
        for item in rows:
            value = (item.get(field) or "").strip()
            if value and value.lower() not in allowed:
                out.append({"egg_id": item.get("egg_id"), "field": field,
                            "value": value,
                            "reason": "khong co trong danh muc bang thong ke"})
    return out
