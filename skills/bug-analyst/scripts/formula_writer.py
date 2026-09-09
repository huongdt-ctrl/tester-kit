#!/usr/bin/env python3
"""Sinh lai toan bo cong thuc thong ke cua sheet 'List bug'.

Skill KHONG dich cong thuc cu cua template -- no viet lai tu dau theo Layout.
Ly do: template goc co san 2 loi cong thuc, dich cong thuc sai chi tao ra cong
thuc sai o toa do moi.

Hai loi cua template duoc sua o day:
  - Section III (Milestone) dem $G$ (Priority) thay vi $E$ (Milestone), va chia
    phan tram cho $B$76 (TOTAL cua Priority) thay vi TOTAL cua chinh no.
  - Section VII (Root cause) viet COUNTIF($K$3:$K$39, $A$126:$A$151) -- tham so
    thu hai la ca dai o, khien moi dong tra ve cung mot so.
"""
import re
from copy import copy

from openpyxl.utils import get_column_letter

from template_layout import DATA_START, ISSUE_TABLE_SLOTS

PERCENT_FORMAT = "0.0%"


def clear_data_region(ws, layout, last_col=13):
    """Xoa du lieu mau cua template. Style + dropdown giu nguyen."""
    for row in layout.data_rows:
        for col in range(1, last_col + 1):
            ws.cell(row=row, column=col).value = None


def write_data_rows(ws, layout, rows):
    """Do danh sach bug da map vao vung data. rows: list dict theo key cot."""
    from template_layout import (COL_CAUSE, COL_EGG_ID, COL_EGG_NAME,
                                 COL_MILESTONE, COL_NO, COL_NOTE, COL_PRIORITY,
                                 COL_ROOT_CAUSE, COL_SCREEN, COL_SEVERITY,
                                 COL_TITLE, COL_TYPE)
    mapping = (
        (COL_NO, "no"), (COL_EGG_ID, "egg_id"), (COL_EGG_NAME, "egg_name"),
        (COL_SCREEN, "screen"), (COL_MILESTONE, "milestone"), (COL_TITLE, "title"),
        (COL_PRIORITY, "priority"), (COL_SEVERITY, "severity"), (COL_TYPE, "bug_type"),
        (COL_CAUSE, "cause"), (COL_ROOT_CAUSE, "root_cause"), (COL_NOTE, "note"),
    )
    capacity = layout.data_end - DATA_START + 1
    if len(rows) > capacity:
        # Ghi qua vung data la de len bang thong ke ben duoi -> hong file ma
        # khong bao gi. Chan o day de thanh loi ro rang.
        raise ValueError("Vung data chi chua duoc %d dong nhung co %d bug -- "
                         "Layout tinh thieu dong" % (capacity, len(rows)))
    for idx, item in enumerate(rows):
        row = DATA_START + idx
        for col, key in mapping:
            value = normalize_cell(item.get(key))
            ws["%s%d" % (col, row)] = value


def normalize_cell(value):
    """Cat dau cach dau/cuoi cua moi gia tri chu truoc khi ghi.

    BAT BUOC, khong phai lam dep: nhan cua bang thong ke duoc sinh tu chinh du
    lieu nay va da bi strip. Ghi "Sprint 3 " (con dau cach cuoi) trong khi nhan
    la "Sprint 3" thi COUNTIF cua Excel KHONG khop -- Excel bo qua hoa/thuong
    nhung KHONG bo qua dau cach -> bug do bien mat khoi thong ke, TOTAL im lang
    nho hon so bug.
    """
    if isinstance(value, str):
        value = value.strip()
    return value if value not in ("", None) else None


def write_block(ws, layout, block, labels=None):
    """Viet nhan + cong thuc COUNTIF/percent + dong TOTAL cho mot bang thong ke.

    Nhan duoc ghi cho MOI block, khong chi block dong: bang danh muc co dinh
    cung can ghi lai vi co the vua duoc noi them dong cho nhung muc ma template
    bo sot (vd 5 muc "... Other" cua bang Cause).
    """
    first, last, total = block["first"], block["last"], block["total"]
    label_col, count_col = block["label_col"], block["count_col"]
    pct_col, src_col = block["pct_col"], block["src_col"]
    ds, de = DATA_START, layout.data_end

    if labels is not None:
        labels = list(labels)
        for offset in range(block["slots"]):
            row = first + offset
            ws["%s%d" % (label_col, row)] = labels[offset] if offset < len(labels) else None

    for row in range(first, last + 1):
        ws["%s%d" % (count_col, row)] = (
            "=COUNTIF($%s$%d:$%s$%d,$%s%d)" % (src_col, ds, src_col, de, label_col, row))
        pct = ws["%s%d" % (pct_col, row)]
        # Chia cho TOTAL cua CHINH block nay -- template goc chia nham block khac.
        pct.value = "=IF($%s$%d=0,0,%s%d/$%s$%d)" % (
            count_col, total, count_col, row, count_col, total)
        pct.number_format = PERCENT_FORMAT

    ws["%s%d" % (label_col, total)] = "TOTAL"
    ws["%s%d" % (count_col, total)] = "=SUM(%s%d:%s%d)" % (count_col, first, count_col, last)
    total_pct = ws["%s%d" % (pct_col, total)]
    total_pct.value = "=SUM(%s%d:%s%d)" % (pct_col, first, pct_col, last)
    total_pct.number_format = PERCENT_FORMAT


def read_fixed_labels(ws, block):
    """Nhan danh muc co dinh (priority/severity/cause/root cause) doc thang tu
    template -- khong hardcode lai, tranh lech vi template co o dat 2 dau cach."""
    out = []
    for row in range(block["first"], block["last"] + 1):
        value = ws["%s%d" % (block["label_col"], row)].value
        out.append(str(value).strip() if value is not None else None)
    return out


ALREADY_NUMBERED = re.compile(r"^(\d+[.)]|[-*\u2022])\s")
LINE_HEIGHT = 15.0
SUB_PREFIX = "   \u2022 "          # thut vao qua "1. " roi moi den y con
DEFAULT_COL_WIDTH = 8.43           # width Excel mac dinh khi column_dimensions trong


def _as_blocks(value):
    """Gia tri o -> danh sach BLOCK; moi block = [dong chinh, y con, y con...].

    Nhan:
      - string            -> tach theo newline, moi dong mot block 1 dong
      - list[str]         -> moi phan tu mot block 1 dong
      - list[list[str]]   -> phan tu dau la dong chinh, con lai la y con

    Cat dau cach va bo dong rong: dong rong lot vao giua lam Excel gian chieu
    cao vo ich.
    """
    if value is None:
        return []
    items = value if isinstance(value, (list, tuple)) else str(value).splitlines()
    blocks = []
    for item in items:
        if isinstance(item, (list, tuple)):
            lines = [str(x).strip() for x in item if str(x).strip()]
        else:
            lines = [str(item).strip()] if str(item).strip() else []
        if lines:
            blocks.append(lines)
    return blocks


def _render_cell(blocks):
    """Block -> text cua o: dong chinh danh so, y con thut vao mot gach dau dong.

    Ly do tach y con ra dong rieng: nhoi "viec + bang chung + han" vao cung mot
    dong thi doc phai do mat tim dau la vat, va dong do dai gap 2-3 lan be rong
    cot nen Excel wrap tuy y giua cau.
    """
    if not blocks:
        return None
    if len(blocks) == 1 and len(blocks[0]) == 1:
        return blocks[0][0]
    # Chi danh so khi co TU HAI block: o Issue la mot phat bieu (van de + co che),
    # danh so "1./2." vao mot phat bieu doc ra nhu danh sach viec phai lam.
    numbered = (len(blocks) > 1
                and not all(ALREADY_NUMBERED.match(b[0]) for b in blocks))
    out = []
    for idx, block in enumerate(blocks, 1):
        out.append("%d. %s" % (idx, block[0]) if numbered else block[0])
        out.extend(SUB_PREFIX + line for line in block[1:])
    return "\n".join(out)


def _cell_capacity(ws, coord):
    """So ky tu mot dong hien thi duoc trong o -- CONG be rong ca vung merge.

    O `Issue` la merge B:D, `Action` la merge E:I: tinh theo mot cot thi hut
    khoang 2-5 lan, chieu cao dong ra thieu va text bi che.
    """
    cell = ws[coord]
    cols = [cell.column_letter]
    for rng in ws.merged_cells.ranges:
        if (rng.min_row <= cell.row <= rng.max_row
                and rng.min_col <= cell.column <= rng.max_col):
            # get_column_letter, KHONG dung cell.column_letter: o trong vung
            # merge la MergedCell va MergedCell khong co attribute do.
            cols = [get_column_letter(c)
                    for c in range(rng.min_col, rng.max_col + 1)]
            break
    total = sum(ws.column_dimensions[c].width or DEFAULT_COL_WIDTH for c in cols)
    return max(10, int(total))


def _visual_lines(text, capacity):
    """So dong Excel THAT SU ve ra sau khi wrap, khong phai so dong logic.

    Do that: cot Action rong ~82 ky tu nhung dong action dai 150-200 ky tu ->
    moi dong logic an 2-3 dong hien thi. Tinh height theo so dong logic la dat
    chieu cao thieu mot nua va text bi che im lang.
    """
    if not text:
        return 0
    return sum(max(1, -(-len(line) // capacity)) for line in str(text).split("\n"))


def _align_left(cell):
    """Doi rieng horizontal -> left, giu nguyen vertical / wrap cua template.

    Template can GIUA (`horizontal=center`) ca o Issue, hop voi nhan ngan nhung
    doan van dai thi lech mep hai ben, doc rat met. Phai copy Alignment cu roi
    doi mot field: gan Alignment moi tay se mat `vertical=center` + `wrap_text`,
    ma mat wrap_text la text nhieu dong tran ngang qua o ben canh.
    """
    alignment = copy(cell.alignment)
    alignment.horizontal = "left"
    cell.alignment = alignment


def write_issue_actions(ws, layout, issues):
    """Bang VIII - Issue / Action / Status / PIC (cot A,B,E,J,K theo template).

    Noi dung do agent viet (xem SKILL.md muc 7b). Ham nay lo hai thu ma viec ghi
    tay hay quen:

    - **Chieu cao dong**: o B/E la merged cell + wrap_text, va Excel KHONG tu
      gian chieu cao dong da merge -> text nhieu dong bi che mat neu khong set
      height. Tinh theo so dong SAU KHI WRAP, khong phai so dong logic.
    - **Tran 12 dong** (row 155..166): duoi 166 khong con o nao co border/merge/
      wrap, ghi tran xuong do la du lieu nam tren vung trang ma khong ai thay.
      Vuot tran thi bao ra `issues_dropped`, khong cat im lang.
    """
    header = layout.issue_table_header
    written, dropped = 0, []
    for idx, item in enumerate(issues or []):
        issue_text = _render_cell(_as_blocks(item.get("issue")))
        action_text = _render_cell(_as_blocks(item.get("action")))
        if idx >= ISSUE_TABLE_SLOTS:
            dropped.append((issue_text or "(issue rong)").split("\n")[0])
            continue
        row = header + 1 + idx
        ws["A%d" % row] = idx + 1
        ws["B%d" % row] = issue_text
        ws["E%d" % row] = action_text
        # O B la merged cell B:D -> style cua o goc quyet dinh ca vung merge.
        _align_left(ws["B%d" % row])
        ws["J%d" % row] = item.get("status") or "Open"
        ws["K%d" % row] = item.get("pic")
        n_lines = max(_visual_lines(issue_text, _cell_capacity(ws, "B%d" % row)),
                      _visual_lines(action_text, _cell_capacity(ws, "E%d" % row)))
        if n_lines > 1:
            ws.row_dimensions[row].height = n_lines * LINE_HEIGHT
        written += 1
    return {"issues_written": written, "issues_dropped": dropped,
            "issue_slots": ISSUE_TABLE_SLOTS}


def write_phase_header(ws, phase, date_from, date_to):
    """Ghi giai doan + khoang ngay thu bug vao tieu de section I.

    Ca dong 1 la mot o gop A1:L1 -> chi ghi duoc vao o goc A1, ghi vao F1 se
    nem AttributeError vi MergedCell la read-only.
    """
    label = "I - CLASSIFICATION OF BUG — Giai đoạn: %s | Bug tạo từ %s đến %s" % (
        phase or "-", date_from or "-", date_to or "-")
    ws["A1"] = label
    return label
