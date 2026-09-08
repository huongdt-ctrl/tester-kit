#!/usr/bin/env python3
"""Chen dong vao sheet va va lai 4 thu ma openpyxl KHONG tu va.

Da do thuc te tren template TEM-ST03_02: ws.insert_rows() day gia tri + style
cua o xuong dung cho, nhung BO LAI:
  1. merged cell range   -> van o toa do cu
  2. data validation sqref -> van G3:G39
  3. chart series ref    -> van tro dong cu
  4. chart anchor        -> chart nam de len bang khac
Bo qua bat ky muc nao trong 4 muc tren = file mo ra trong Excel bi lech am tham.
Module nay ton tai chi de dam bao ca 4 deu duoc va lai.
"""
import re
from copy import copy

from openpyxl.utils import range_boundaries
from openpyxl.worksheet.cell_range import MultiCellRange

from template_layout import (BLOCK_BY_KEY, DATA_LAST_COL, DATA_START,
                             VALIDATED_COLS)

REF_RE = re.compile(r"^(?:'[^']+'|[^!]+)!\$([A-Z]+)\$(\d+)(?::\$([A-Z]+)\$(\d+))?$")


def _parse_ref(ref):
    """'List bug'!$A$105:$A$121 -> ('A', 105, 'A', 121). None neu khong parse duoc."""
    m = REF_RE.match(str(ref or "").strip())
    if not m:
        return None
    col1, row1, col2, row2 = m.groups()
    return col1, int(row1), col2 or col1, int(row2 or row1)


def _ref(sheet_name, col, row1, row2=None):
    if row2 is None:
        return "'%s'!$%s$%d" % (sheet_name, col, row1)
    return "'%s'!$%s$%d:$%s$%d" % (sheet_name, col, row1, col, row2)


def copy_row_style(ws, src_row, dst_row, last_col=DATA_LAST_COL):
    """Nhan style tu mot dong mau sang dong moi chen (openpyxl chen ra dong tran)."""
    for col in range(1, last_col + 1):
        src = ws.cell(row=src_row, column=col)
        dst = ws.cell(row=dst_row, column=col)
        dst._style = copy(src._style)


def apply_insertions(ws, layout):
    """Chen dong theo ke hoach cua Layout, tu duoi len, va va lai merged range."""
    if not layout.insertions:
        return
    base_merges = [str(rng) for rng in ws.merged_cells.ranges]
    for pos, count in layout.insertions_bottom_up():
        ws.insert_rows(pos, count)
        # pos-1 la dong cuoi cua khoi vua duoc noi dai -> lay lam mau style.
        for offset in range(count):
            copy_row_style(ws, pos - 1, pos + offset)
    _remap_merges(ws, layout, base_merges)
    _replicate_new_row_merges(ws, layout)


def _replicate_new_row_merges(ws, layout):
    """Dong moi chen KHONG duoc thua huong o gop cua dong mau.

    Bang Cause/Root cause gop A:C tren TUNG dong; dong moi khong gop thi nhan
    bi cat o cot A trong khi cac dong khac trai het A:C -> nhin la thay lech.
    """
    for pos, count in layout.insertions:
        src = layout.shift(pos - 1)
        spans = [(rng.min_col, rng.max_col) for rng in ws.merged_cells.ranges
                 if rng.min_row == rng.max_row == src]
        for offset in range(1, count + 1):
            for min_col, max_col in spans:
                ws.merge_cells(start_row=src + offset, end_row=src + offset,
                               start_column=min_col, end_column=max_col)


def _remap_merges(ws, layout, base_merges):
    for rng in base_merges:
        try:
            ws.unmerge_cells(rng)
        except (KeyError, ValueError):
            pass
    for rng in base_merges:
        try:
            col1, row1, col2, row2 = _split_plain(rng)
        except ValueError:
            continue
        ws.merge_cells(start_row=layout.shift(row1), end_row=layout.shift(row2),
                       start_column=col1, end_column=col2)


def _split_plain(rng):
    """'A105:C105' -> (col1_idx, 105, col3_idx, 105)."""
    col1, row1, col2, row2 = range_boundaries(rng)
    if None in (col1, row1, col2, row2):
        raise ValueError(rng)
    return col1, row1, col2, row2


def remap_validations(ws, layout):
    """Keo sqref cua moi dropdown phu het vung data that (row 3..data_end)."""
    touched = []
    for dv in ws.data_validations.dataValidation:
        parsed = _split_first_cell(str(dv.sqref))
        if parsed is None:
            continue
        col_letter = parsed
        if col_letter not in VALIDATED_COLS:
            continue
        dv.sqref = MultiCellRange("%s%d:%s%d" % (col_letter, DATA_START,
                                                 col_letter, layout.data_end))
        touched.append(col_letter)
    return touched


def _split_first_cell(sqref):
    m = re.match(r"^([A-Z]+)\d+", str(sqref).strip())
    return m.group(1) if m else None


def remap_charts(ws, layout, sheet_name):
    """Tro lai 5 chart vao dung bang cua no, va sua luon 4 loi ref cua template.

    Chart duoc nhan dien bang dong bat dau cua ref GOC (sai so <= 2 dong), khong
    dua vao thu tu ws._charts -- thu tu do khong co gi bao dam.
    """
    fixes = []
    blocks = layout.blocks()
    for chart in list(ws._charts):
        block = _match_block(chart, blocks)
        if block is None:
            continue
        for ser in chart.series:
            _retarget_series(ser, block, sheet_name, fixes)
        _shift_anchor(chart, layout)
    return fixes


def _match_block(chart, blocks):
    for ser in chart.series:
        ref = None
        if ser.cat is not None:
            ref = (ser.cat.numRef.f if ser.cat.numRef else
                   (ser.cat.strRef.f if ser.cat.strRef else None))
        parsed = _parse_ref(ref)
        if not parsed:
            continue
        start_row = parsed[1]
        for block in blocks:
            # Ref trong template dung toa do GOC -> so voi first GOC, chua dich.
            if abs(start_row - _base_first(block)) <= 2:
                return block
    return None


def _base_first(block):
    """first cua block o toa do GOC (ref trong template chua bi dich)."""
    return BLOCK_BY_KEY[block["key"]].first


def _retarget_series(ser, block, sheet_name, fixes):
    first, last = block["first"], block["last"]
    cat_ref = _ref(sheet_name, block["label_col"], first, last)
    val_ref = _ref(sheet_name, block["count_col"], first, last)
    if ser.cat is not None and ser.cat.numRef is not None:
        _note(fixes, block, "cat", ser.cat.numRef.f, cat_ref)
        ser.cat.numRef.f = cat_ref
    elif ser.cat is not None and ser.cat.strRef is not None:
        _note(fixes, block, "cat", ser.cat.strRef.f, cat_ref)
        ser.cat.strRef.f = cat_ref
    if ser.val is not None and ser.val.numRef is not None:
        _note(fixes, block, "val", ser.val.numRef.f, val_ref)
        ser.val.numRef.f = val_ref
    # Series title: tro ve o header cua cot dem, thay vi mot o du lieu.
    if ser.tx is not None and ser.tx.strRef is not None:
        head_ref = _ref(sheet_name, block["count_col"], block["header"])
        _note(fixes, block, "tx", ser.tx.strRef.f, head_ref)
        ser.tx.strRef.f = head_ref


def _note(fixes, block, kind, old, new):
    if str(old) != str(new):
        fixes.append({"block": block["key"], "part": kind,
                      "from": str(old), "to": new})


def _shift_anchor(chart, layout):
    anchor = getattr(chart, "anchor", None)
    for side in ("_from", "to"):
        marker = getattr(anchor, side, None)
        if marker is None or not hasattr(marker, "row"):
            continue
        # marker.row la 0-based; Layout.shift lam viec tren so dong 1-based.
        marker.row = layout.shift(marker.row + 1) - 1
