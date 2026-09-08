#!/usr/bin/env python3
"""Test: build file phan tich bug vuot suc chua template (>37 bug, >12 man hinh).

Trong tam: sau khi chen dong, 4 thu ma openpyxl khong tu va (merged range,
dropdown sqref, chart ref, chart anchor) co thuc su duoc va lai khong.
"""
import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "scripts"))

import openpyxl

import workbook_builder as builder
from template_layout import SHEET_NAME, Layout

TEMPLATE = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                        "..", "templates", "template_bug_analysis.xlsx")


def make_rows(count, screens=4, milestones=2):
    return [{
        "no": i + 1, "egg_id": "#%d" % (100 + i), "egg_name": "Bug so %d" % (i + 1),
        "screen": "G%d" % (i % screens), "milestone": "Sprint %d" % (i % milestones + 1),
        "title": "[EPIC-1][G%d] loi %d" % (i % screens, i + 1),
        "priority": "High", "severity": "Major", "bug_type": "Logic",
        "cause": "COD1.2 Coding Logic Mistake",
        "root_cause": "SKI3.3 Inadequate skill in Testing", "note": "",
    } for i in range(count)]


class TestBuilderExpansion(unittest.TestCase):
    def setUp(self):
        if not os.path.exists(TEMPLATE):
            self.skipTest("khong tim thay template goc")
        self.tmp = tempfile.mkdtemp()

    def _build(self, rows):
        out = os.path.join(self.tmp, "out.xlsx")
        report = builder.build(TEMPLATE, out, rows, phase="Sprint 3",
                               date_from="2026-08-01", date_to="2026-08-31",
                               overwrite=True)
        return out, report

    def test_builder_within_capacity_keeps_coordinates_above_cause_table(self):
        # Bang Cause LUON duoc noi them 5 dong (template bo sot 5 muc danh muc),
        # nen moi toa do TREN bang do phai giu nguyen, duoi thi dich.
        # Arrange
        rows = make_rows(10, screens=4, milestones=2)
        # Act
        out, report = self._build(rows)
        # Assert
        self.assertEqual(
            {k: v for k, v in report["rows_inserted"].items() if v},
            {"cause": 5})
        ws = openpyxl.load_workbook(out)[SHEET_NAME]
        self.assertEqual(ws["A55"].value, "TOTAL")      # TOTAL bang man hinh
        self.assertEqual(ws["B3"].value, "#100")

    def test_builder_completes_cause_catalogue_missing_from_template(self):
        # Arrange
        rows = make_rows(4)
        # Act
        out, report = self._build(rows)
        # Assert -- 5 muc "... Other" duoc noi vao bang VI
        self.assertEqual(report["catalogue_rows_added"]["cause"],
                         ["REQ1.4 Other", "COD1.6 Other", "TES1.3 Other",
                          "DEP1.3 Other", "OTH Other"])
        self.assertNotIn("rootcause", report["catalogue_rows_added"])
        layout = Layout(4, {"cause": 22, "rootcause": 26})
        block = layout.block("cause")
        self.assertEqual(block["slots"], 22)
        ws = openpyxl.load_workbook(out)[SHEET_NAME]
        labels = [ws["A%d" % r].value for r in range(block["first"], block["last"] + 1)]
        self.assertEqual(labels[-1], "OTH Other")
        self.assertEqual(ws["A%d" % block["total"]].value, "TOTAL")

    def test_builder_counts_bug_classified_into_completed_catalogue_row(self):
        # Truoc khi noi danh muc, bug xep vao "OTH Other" khong duoc dem vao dau.
        # Arrange
        rows = make_rows(3)
        rows[0]["cause"] = "OTH Other"
        # Act
        out, report = self._build(rows)
        # Assert
        self.assertEqual(report["uncounted"], [])
        layout = Layout(3, {"cause": 22, "rootcause": 26})
        block = layout.block("cause")
        ws = openpyxl.load_workbook(out)[SHEET_NAME]
        row = next(r for r in range(block["first"], block["last"] + 1)
                   if ws["A%d" % r].value == "OTH Other")
        self.assertEqual(ws["D%d" % row].value, "=COUNTIF($J$3:$J$39,$A%d)" % row)

    def test_builder_over_capacity_shifts_every_dependent_coordinate(self):
        # Arrange: 50 bug (>37), 20 man hinh (>12), 8 milestone (>5)
        rows = make_rows(50, screens=20, milestones=8)
        # Act
        out, report = self._build(rows)
        # Assert
        self.assertEqual(report["bug_count"], 50)
        self.assertEqual({k: v for k, v in report["rows_inserted"].items() if v},
                         {"data": 13, "screen": 8, "milestone": 3, "cause": 5})
        layout = Layout(50, {"screen": 20, "milestone": 8, "cause": 22, "rootcause": 26})
        ws = openpyxl.load_workbook(out)[SHEET_NAME]

        # du lieu bug ghi het, khong bi cat o dong 39
        self.assertEqual(ws["B%d" % layout.data_end].value, "#149")
        # TOTAL cua tung block nam dung dong moi
        for key in ("screen", "milestone", "priority", "severity", "cause", "rootcause"):
            block = layout.block(key)
            self.assertEqual(ws["%s%d" % (block["label_col"], block["total"])].value,
                             "TOTAL", "TOTAL sai cho o block %s" % key)
        # dropdown phu het vung data moi
        sqrefs = {str(dv.sqref) for dv in ws.data_validations.dataValidation}
        self.assertIn("G3:G%d" % layout.data_end, sqrefs)
        self.assertIn("K3:K%d" % layout.data_end, sqrefs)

    def test_builder_over_capacity_repoints_all_five_charts(self):
        # Arrange
        rows = make_rows(50, screens=20, milestones=8)
        # Act
        out, _ = self._build(rows)
        # Assert
        layout = Layout(50, {"screen": 20, "milestone": 8, "cause": 22, "rootcause": 26})
        ws = openpyxl.load_workbook(out)[SHEET_NAME]
        self.assertEqual(len(ws._charts), 5)
        expect = set()
        for key in ("screen", "priority", "severity", "cause", "rootcause"):
            b = layout.block(key)
            expect.add("'%s'!$%s$%d:$%s$%d" % (SHEET_NAME, b["count_col"], b["first"],
                                               b["count_col"], b["last"]))
        actual = {s.val.numRef.f for ch in ws._charts for s in ch.series
                  if s.val and s.val.numRef}
        self.assertEqual(actual, expect)

    def test_builder_charts_read_count_column_not_stale_column_b(self):
        # Template goc: chart VI/VII tro $B$ nhung so lieu nam o cot D.
        # Arrange
        rows = make_rows(5)
        # Act
        out, report = self._build(rows)
        # Assert
        ws = openpyxl.load_workbook(out)[SHEET_NAME]
        refs = {s.val.numRef.f for ch in ws._charts for s in ch.series
                if s.val and s.val.numRef}
        layout = Layout(5, {"cause": 22, "rootcause": 26})
        for key in ("cause", "rootcause"):
            b = layout.block(key)
            self.assertIn("'%s'!$D$%d:$D$%d" % (SHEET_NAME, b["first"], b["last"]), refs)
        self.assertTrue(report["chart_refs_fixed"])

    def test_builder_milestone_section_counts_milestone_column_not_priority(self):
        # Arrange
        rows = make_rows(6, screens=2, milestones=3)
        # Act
        out, _ = self._build(rows)
        # Assert
        ws = openpyxl.load_workbook(out)[SHEET_NAME]
        self.assertIn("$E$3:$E$39", ws["B61"].value)
        self.assertNotIn("$G$", ws["B61"].value)

    def test_builder_rootcause_countif_uses_single_label_cell(self):
        # Arrange
        rows = make_rows(6)
        # Act
        out, _ = self._build(rows)
        # Assert
        block = Layout(6, {"cause": 22, "rootcause": 26}).block("rootcause")
        self.assertEqual(
            openpyxl.load_workbook(out)[SHEET_NAME]["D%d" % block["first"]].value,
            "=COUNTIF($K$3:$K$39,$A%d)" % block["first"])

    def test_builder_reports_values_outside_statistics_catalogue(self):
        # Gia tri khong co trong CA danh muc (vd go tay sai) van phai bao ra.
        # Arrange
        rows = make_rows(2)
        rows[0]["cause"] = "COD9.9 Go tay sai"
        # Act
        _, report = self._build(rows)
        # Assert
        self.assertEqual(len(report["uncounted"]), 1)
        self.assertEqual(report["uncounted"][0]["value"], "COD9.9 Go tay sai")

    def test_builder_refuses_to_overwrite_existing_file_by_default(self):
        # Arrange
        rows = make_rows(2)
        out, _ = self._build(rows)
        # Act + Assert
        with self.assertRaises(FileExistsError):
            builder.build(TEMPLATE, out, rows, overwrite=False)


if __name__ == "__main__":
    unittest.main(verbosity=2)
