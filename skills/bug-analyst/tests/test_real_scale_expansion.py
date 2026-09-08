#!/usr/bin/env python3
"""Test o dung QUY MO VAN HANH THAT: ~200 bug moi giai doan.

Vi sao can rieng test nay: cac test khac chay 45-50 bug (chen ~8-13 dong).
Muc that la 200 bug -> chen 163 dong vao vung data, day MOI bang thong ke va
ca 5 chart xuong hon 160 dong. Sai lech tich luy chi lo ra o quy mo nay.

Bat bien duy nhat can giu: MOI bang phai dem dung 200, khong bang nao thieu
hay doi.
"""
import os
import sys
import tempfile
import unittest
from collections import Counter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import openpyxl

import workbook_builder as builder
from helpers import TEMPLATE
from template_layout import SHEET_NAME, Layout

BUG_COUNT = 200
# 18 man hinh (>12 slot template) va 7 milestone (>5 slot) -> ca hai bang deu gian.
SCREENS = ["G10", "G10.1", "G10.4", "G12 13 14", "G16.1", "G16.2", "G17", "G17.2",
           "G20", "M03", "M07", "M08", "M09", "All screen", "G21", "G22", "G23", "G24"]
MILESTONES = ["Sprint 1", "Sprint 2", "Sprint 3", "Sprint 4", "Sprint 5", "Sprint 6", "UAT"]
PRIORITIES = ["Emergent", "Urgent", "High", "Normal", "Low"]
SEVERITIES = ["Fatal", "Critical", "Major", "Minor", "Low"]
# Gom ca 2 muc "... Other" ma template goc bo sot khoi bang VI.
CAUSES = ["COD1.2 Coding Logic Mistake", "REQ1.2 Unclear Definition", "OTH Other",
          "DES1.4 Missing or Incomplete  UI/UX Design", "TES1.1 Test case missing",
          "DEP1.3 Other"]
ROOTS = ["SKI3.3 Inadequate skill in Testing", "PRO2.3 M-I Review-Testing Process",
         "IMP3 Carelessness"]

ALL_BLOCKS = ("screen", "milestone", "priority", "severity", "cause", "rootcause")


def make_rows(count):
    return [{
        "no": i + 1, "egg_id": "#%d" % (1000 + i), "egg_name": "Loi %d" % (i + 1),
        "screen": SCREENS[i % len(SCREENS)], "milestone": MILESTONES[i % len(MILESTONES)],
        "title": "[GET-1][%s] loi %d" % (SCREENS[i % len(SCREENS)], i + 1),
        "priority": PRIORITIES[i % len(PRIORITIES)],
        "severity": SEVERITIES[i % len(SEVERITIES)],
        "bug_type": "UI" if i % 2 else "Logic",
        "cause": CAUSES[i % len(CAUSES)], "root_cause": ROOTS[i % len(ROOTS)],
        "note": "https://gl.test/-/issues/%d" % (1000 + i),
    } for i in range(count)]


def countif_total(labels, cells):
    """SUM(COUNTIF(...)) theo luat Excel: casefold, giu dau cach."""
    counter = Counter(str(v).lower() for v in cells if v not in (None, ""))
    return sum(counter[str(lb).lower()] for lb in labels if lb)


class TestRealScaleExpansion(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if not os.path.exists(TEMPLATE):
            raise unittest.SkipTest("khong tim thay template goc")
        cls.tmp = tempfile.mkdtemp()
        cls.out = os.path.join(cls.tmp, "scale.xlsx")
        cls.report = builder.build(
            TEMPLATE, cls.out, make_rows(BUG_COUNT), phase="Sprint 6",
            date_from="2026-08-01", date_to="2026-08-31",
            issues=[{"issue": "Bug UI lap lai nhieu sprint",
                     "action": "Bo sung checklist UI", "pic": "huongdt"}],
            overwrite=True)
        cls.ws = openpyxl.load_workbook(cls.out)[SHEET_NAME]
        cls.layout = Layout(BUG_COUNT, {"screen": len(cls.report["screens"]),
                                        "milestone": len(cls.report["milestones"]),
                                        "cause": 22, "rootcause": 26})

    def _block_cells(self, key):
        block = self.layout.block(key)
        labels = [self.ws["%s%d" % (block["label_col"], r)].value
                  for r in range(block["first"], block["last"] + 1)]
        cells = [self.ws["%s%d" % (block["src_col"], r)].value
                 for r in self.layout.data_rows]
        return block, labels, cells

    def test_scale_every_statistics_table_counts_all_two_hundred_bugs(self):
        # Arrange + Act da xong o setUpClass
        # Assert
        for key in ALL_BLOCKS:
            _, labels, cells = self._block_cells(key)
            self.assertEqual(countif_total(labels, cells), BUG_COUNT,
                             "bang %s dem sai o quy mo 200 bug" % key)

    def test_scale_all_two_hundred_bugs_written_without_truncation(self):
        # Assert
        self.assertEqual(self.report["bug_count"], BUG_COUNT)
        self.assertEqual(self.layout.data_end, 2 + BUG_COUNT)
        self.assertEqual(self.ws["B%d" % self.layout.data_end].value,
                         "#%d" % (1000 + BUG_COUNT - 1))

    def test_scale_dropdowns_cover_the_whole_expanded_data_region(self):
        # Assert
        sqrefs = {str(dv.sqref) for dv in self.ws.data_validations.dataValidation}
        for col in ("G", "H", "I", "J", "K"):
            self.assertIn("%s3:%s%d" % (col, col, self.layout.data_end), sqrefs)

    def test_scale_all_five_charts_follow_their_table_down(self):
        # Assert
        expected = set()
        for key in ("screen", "priority", "severity", "cause", "rootcause"):
            b = self.layout.block(key)
            expected.add("'%s'!$%s$%d:$%s$%d" % (SHEET_NAME, b["count_col"], b["first"],
                                                 b["count_col"], b["last"]))
        actual = {s.val.numRef.f for ch in self.ws._charts for s in ch.series
                  if s.val and s.val.numRef}
        self.assertEqual(actual, expected)

    def test_scale_tables_stay_in_order_without_overlapping(self):
        # Bang nao de len bang khac la mat du lieu thong ke ma khong ai thay.
        # Assert
        edges = [(k, self.layout.block(k)) for k in ALL_BLOCKS]
        for (key, block) in edges:
            self.assertEqual(block["last"] + 1, block["total"],
                             "bang %s co dong TOTAL khong lien tiep" % key)
        for (key_a, a), (key_b, b) in zip(edges, edges[1:]):
            self.assertLess(a["total"], b["header"],
                            "bang %s de len bang %s" % (key_a, key_b))
        last = edges[-1][1]
        self.assertGreater(self.layout.issue_table_header, last["total"],
                           "bang Issue/Action de len bang Root cause")
        self.assertEqual(self.ws["A%d" % self.layout.issue_table_header].value, "No.")


if __name__ == "__main__":
    unittest.main(verbosity=2)
