#!/usr/bin/env python3
"""Test: bang Issue / Action -- dinh dang o va tran 12 dong.

Bang nay do agent viet tay nen script phai lo hai thu agent hay quen: xuong dong
+ danh so cho de doc, va chieu cao dong (o B/E la merged cell, Excel KHONG tu
gian chieu cao dong da merge -> text nhieu dong bi che mat). Xem SKILL.md muc 7b.
"""
import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "scripts"))

import openpyxl

import formula_writer
import workbook_builder as builder
from template_layout import SHEET_NAME, ISSUE_TABLE_SLOTS, Layout

TEMPLATE = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                        "..", "templates", "template_bug_analysis.xlsx")


def one_row():
    return [{"no": 1, "egg_id": "#500", "egg_name": "Loi 1", "screen": "G10.1",
             "milestone": "Sprint 3", "title": "[GET-1][G10.1] loi 1",
             "priority": "High", "severity": "Major", "bug_type": "UI",
             "cause": "COD1.2 Coding Logic Mistake", "root_cause": "IMP3 Carelessness",
             "note": "https://gl.test/-/issues/500"}]


class IssueActionTableTest(unittest.TestCase):

    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.layout = Layout(1, {"cause": 22, "rootcause": 26})
        self.header = self.layout.issue_table_header

    def build(self, issues, name="issues.xlsx"):
        out = os.path.join(self.tmp, name)
        report = builder.build(TEMPLATE, out, one_row(), phase="Sprint 3",
                               date_from="2026-08-01", date_to="2026-08-31",
                               issues=issues, overwrite=True)
        return report, openpyxl.load_workbook(out)[SHEET_NAME]

    def test_issue_action_list_value_is_numbered_one_per_line(self):
        # Arrange -- action nhieu viec, truyen dang list
        issues = [{"issue": "12/45 bug o Cause = COD1.1",
                   "action": ["Nam (BE) them unit test -- bang chung: link MR",
                              "Linh (QA) them 6 test case bien -- file v1.2"],
                   "pic": "Nam"}]
        # Act
        report, ws = self.build(issues)
        # Assert
        self.assertEqual(
            ws["E%d" % (self.header + 1)].value,
            "1. Nam (BE) them unit test -- bang chung: link MR\n"
            "2. Linh (QA) them 6 test case bien -- file v1.2")
        self.assertEqual(report["issues_written"], 1)

    def test_issue_action_pre_numbered_lines_are_not_numbered_twice(self):
        # Arrange -- agent tu danh so san trong string
        issues = [{"issue": "8 bug man Dang nhap",
                   "action": "1. Nam sua validate\n2. Linh bo sung test case"}]
        # Act
        _, ws = self.build(issues, "prenumbered.xlsx")
        # Assert
        self.assertEqual(ws["E%d" % (self.header + 1)].value,
                         "1. Nam sua validate\n2. Linh bo sung test case")

    def test_issue_action_single_line_keeps_template_row_height(self):
        # Arrange
        issues = [{"issue": "8 bug man Dang nhap", "action": "Nam sua validate"}]
        # Act
        _, ws = self.build(issues, "single.xlsx")
        # Assert -- mot dong thi khong dung vao height cua template
        self.assertIsNone(ws.row_dimensions[self.header + 1].height)

    def test_issue_action_multiline_sets_row_height_because_merged_cell_never_autofits(self):
        # Arrange -- 3 dong action
        issues = [{"issue": "12/45 bug o Cause = COD1.1",
                   "action": ["viec 1", "viec 2", "viec 3"]}]
        # Act
        _, ws = self.build(issues, "height.xlsx")
        # Assert
        height = ws.row_dimensions[self.header + 1].height
        self.assertIsNotNone(height, "khong set height -> text bi che trong o merged")
        self.assertGreaterEqual(height, 45.0)

    def test_issue_action_over_twelve_entries_reports_dropped_instead_of_overflowing(self):
        # Arrange -- 14 issue, template chi co 12 o co border/merge/wrap
        issues = [{"issue": "issue %d" % i, "action": "viec %d" % i}
                  for i in range(1, 15)]
        # Act
        report, ws = self.build(issues, "overflow.xlsx")
        # Assert
        self.assertEqual(report["issues_written"], ISSUE_TABLE_SLOTS)
        self.assertEqual(report["issues_dropped"], ["issue 13", "issue 14"])
        last = self.header + ISSUE_TABLE_SLOTS
        self.assertEqual(ws["B%d" % last].value, "issue 12")
        self.assertIsNone(ws["B%d" % (last + 1)].value,
                          "tran xuong vung khong co border -> du lieu tang hinh")

    def test_issue_cell_is_left_aligned_but_keeps_template_wrap_and_vertical(self):
        # Template can GIUA o Issue -> doan van dai lech mep hai ben. Doi sang
        # left NHUNG khong duoc lam mat wrap_text (mat wrap la text tran ngang).
        # Arrange
        issues = [{"issue": "11/15 Egg co Cause = COD1.1 Lack in code",
                   "action": ["viec 1", "viec 2"]}]
        # Act
        _, ws = self.build(issues, "align.xlsx")
        # Assert
        cell = ws["B%d" % (self.header + 1)]
        self.assertEqual(cell.alignment.horizontal, "left")
        self.assertTrue(cell.alignment.wrap_text, "mat wrap_text -> text tran ngang")
        self.assertEqual(cell.alignment.vertical, "center")

    def test_issue_action_nested_list_puts_each_sub_idea_on_its_own_line(self):
        # Nhoi "viec + bang chung + han" vao cung mot dong thi doc phai do mat
        # tim dau la vat -> y con phai xuong dong, thut vao qua so.
        # Arrange
        issues = [{"issue": "6/15 Egg root cause IMP3",
                   "action": [["Chan ngay: FE bo hardcode 4 man",
                               "Bang chung: link MR",
                               "Han: 2026-09-15"],
                              ["Quet tuong tu: grep hex tren 30 module",
                               "Bang chung: danh sach file"]]}]
        # Act
        _, ws = self.build(issues, "nested.xlsx")
        # Assert
        self.assertEqual(
            ws["E%d" % (self.header + 1)].value,
            "1. Chan ngay: FE bo hardcode 4 man\n"
            "   \u2022 Bang chung: link MR\n"
            "   \u2022 Han: 2026-09-15\n"
            "2. Quet tuong tu: grep hex tren 30 module\n"
            "   \u2022 Bang chung: danh sach file")

    def test_issue_single_block_is_not_numbered_because_it_is_one_statement(self):
        # O Issue la MOT phat bieu (van de + co che sinh loi). Danh so "1./2."
        # vao do doc ra nhu danh sach viec phai lam -> sai nghia.
        # Arrange
        issues = [{"issue": [["6/15 Egg root cause IMP3, toan bo nhom layout",
                              "Co che: style hardcode nen lech Figma"]],
                   "action": "viec 1"}]
        # Act
        _, ws = self.build(issues, "one_block.xlsx")
        # Assert
        self.assertEqual(ws["B%d" % (self.header + 1)].value,
                         "6/15 Egg root cause IMP3, toan bo nhom layout\n"
                         "   \u2022 Co che: style hardcode nen lech Figma")

    def test_issue_action_row_height_counts_lines_after_wrap_not_logical_lines(self):
        # Cot Action merge E:I rong ~82 ky tu. Mot dong 200 ky tu an 3 dong hien
        # thi -> tinh height theo so dong logic la dat thieu, text bi che.
        # Arrange
        long_line = "x" * 200
        issues = [{"issue": "issue ngan", "action": [long_line]}]
        # Act
        _, ws = self.build(issues, "wrap.xlsx")
        # Assert
        height = ws.row_dimensions[self.header + 1].height
        self.assertGreaterEqual(height, 45.0,
                                "height tinh theo 1 dong logic -> che mat 2 dong")

    def test_issue_action_capacity_sums_whole_merged_range(self):
        # Tinh be rong theo MOT cot thi hut 2-5 lan (Issue merge B:D, Action E:I).
        # Arrange
        _, ws = self.build([{"issue": "a", "action": "b"}], "capacity.xlsx")
        row = self.header + 1
        # Act
        cap_issue = formula_writer._cell_capacity(ws, "B%d" % row)
        cap_action = formula_writer._cell_capacity(ws, "E%d" % row)
        # Assert -- B:D ~37 ky tu, E:I ~82 ky tu
        self.assertGreater(cap_issue, ws.column_dimensions["B"].width)
        self.assertGreater(cap_action, cap_issue)

    def test_issue_action_blank_lines_are_dropped_before_writing(self):
        # Arrange -- dong rong lot vao giua lam gian chieu cao vo ich
        issues = [{"issue": "  8 bug man Dang nhap  ",
                   "action": ["viec 1", "   ", "", "viec 2"]}]
        # Act
        _, ws = self.build(issues, "blank.xlsx")
        # Assert
        self.assertEqual(ws["B%d" % (self.header + 1)].value, "8 bug man Dang nhap")
        self.assertEqual(ws["E%d" % (self.header + 1)].value, "1. viec 1\n2. viec 2")
        self.assertEqual(ws.row_dimensions[self.header + 1].height, 30.0)


if __name__ == "__main__":
    unittest.main()
