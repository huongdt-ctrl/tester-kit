#!/usr/bin/env python3
"""
Local .xlsx backend exposing the same four methods sheet_io calls on a gspread
worksheet: get_all_values, row_values, update_cell, duplicate.

Why an adapter instead of branching inside sheet_io: read_cases / write_result /
assert_writable already encode every column-protection rule. Duplicating that
logic for Excel would mean two places to keep in sync, and the Excel copy would
be the one that quietly loses a guard. With the adapter those functions stay
untouched and source_type only decides which object they receive.

Writes are flushed to disk on every update_cell -- a run interrupted halfway
must not lose the results already recorded, same as Google Sheets behaves.
"""
import datetime
import os
import shutil

from openpyxl import load_workbook


class ExcelWorksheet(object):
    def __init__(self, path, worksheet_name):
        if not os.path.exists(path):
            raise IOError("Khong tim thay file Excel: %s" % path)
        self.path = path
        self.worksheet_name = worksheet_name
        self._wb = load_workbook(path)
        if worksheet_name not in self._wb.sheetnames:
            raise IOError("File %s khong co worksheet %r (co: %s)"
                          % (path, worksheet_name, ", ".join(self._wb.sheetnames)))
        self._ws = self._wb[worksheet_name]

    # --- gspread-compatible surface -------------------------------------
    def get_all_values(self):
        rows = []
        for row in self._ws.iter_rows(values_only=True):
            rows.append(["" if c is None else str(c) for c in row])
        return rows

    def row_values(self, row_number):
        row = next(self._ws.iter_rows(min_row=row_number, max_row=row_number, values_only=True), ())
        return ["" if c is None else str(c) for c in row]

    def update_cell(self, row, col, value):
        self._ws.cell(row=row, column=col, value=value)
        self._wb.save(self.path)          # flush per write, khong giu trong RAM

    def duplicate(self, new_sheet_name=None):
        """Backup as a sibling FILE rather than an extra tab: an .xlsx handed to
        a client should not carry a snapshot tab, and a separate file survives
        the original being overwritten."""
        stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        base, ext = os.path.splitext(self.path)
        target = "%s__%s%s" % (base, new_sheet_name or ("backup_" + stamp), ext)
        shutil.copy2(self.path, target)
        return target
