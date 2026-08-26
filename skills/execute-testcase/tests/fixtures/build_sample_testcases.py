#!/usr/bin/env python3
"""
Build the 5-case draft test case file used by the integration test.

Deliberately NOT five happy-path rows: the mix of Test result values and
Priority/Classification values is what makes the file able to exercise every
branch of case_filter (baseline vs rerun_failed vs selected_cases) and the
duplicate-bug path (TC003 already carries Bug ID 1234).

Regenerate with:  python3 build_sample_testcases.py
"""
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HEADER = ["Classification 1", "TC ID", "Title", "Priority", "Preconditions", "Steps",
          "Test data", "Expected Result", "Test result", "Test date", "Tested by",
          "Remark", "Actual Result", "Bug ID", "Bug URL"]

ROWS = [
    ["Common Test Cases", "TC001", "Mở màn hình Edit Event với event đang Draft", "High",
     "Đã login bằng tài khoản Admin Event; tồn tại 1 event trạng thái Draft",
     "1. Vào menu Event Management\n2. Chọn event trạng thái Draft\n3. Bấm nút Edit",
     "Event: EVT-DRAFT-001",
     "Màn hình Edit Event mở ra, các trường hiển thị đúng dữ liệu hiện tại của event, nút Save và Cancel enable",
     "Untested", "", "", "", "", "", ""],

    ["Functional Test Cases", "TC002", "Lưu event sau khi sửa tên hợp lệ", "High",
     "Đang ở màn hình Edit Event của EVT-DRAFT-001",
     "1. Sửa trường Event Name thành giá trị hợp lệ\n2. Bấm Save\n3. Quay lại danh sách event",
     "Event Name: 'Live Concert Thu Dong 2026'",
     "Hiện message lưu thành công; danh sách event thể hiện tên mới; không phát sinh lỗi console",
     "Untested", "", "", "", "", "", ""],

    ["Functional Test Cases", "TC003", "Không cho lưu khi ngày kết thúc trước ngày bắt đầu", "Medium",
     "Đang ở màn hình Edit Event của EVT-DRAFT-001",
     "1. Đặt Start Date = 2026-09-10\n2. Đặt End Date = 2026-09-01\n3. Bấm Save",
     "Start Date: 2026-09-10 / End Date: 2026-09-01",
     "Hệ thống chặn lưu, hiện message lỗi validate ở trường End Date, dữ liệu event không đổi",
     "Fail", "2026-08-20", "huongdt",
     "Actual khác Expected: hệ thống cho lưu thành công", "Lưu thành công, không hiện message lỗi nào",
     "1234", "https://redmine.example.com/issues/1234"],

    ["Functional Test Cases", "TC004", "Hủy sửa event thì dữ liệu không bị thay đổi", "Low",
     "Đang ở màn hình Edit Event của EVT-DRAFT-001",
     "1. Sửa Event Name\n2. Bấm Cancel\n3. Mở lại màn hình Edit Event",
     "Event Name nhập tạm: 'Tên tạm ABC'",
     "Không lưu thay đổi; Event Name giữ nguyên giá trị trước khi sửa",
     "Pass", "2026-08-20", "huongdt", "", "", "", ""],

    ["Common Test Cases", "TC005", "Sửa event đã Published cần quyền Admin Event", "Medium",
     "Tồn tại 1 event trạng thái Published; tài khoản test có quyền Admin Event",
     "1. Vào Event Management\n2. Chọn event Published\n3. Bấm Edit\n4. Sửa mô tả và Save",
     "Event: EVT-PUB-001",
     "Sửa được mô tả và lưu thành công; lịch sử thay đổi ghi nhận đúng người sửa",
     "Pending", "2026-08-20", "huongdt",
     "Chưa test được: chưa có dữ liệu event ở trạng thái Published", "", "", ""],
]

WIDTHS = {"A": 22, "B": 9, "C": 42, "D": 9, "E": 38, "F": 38, "G": 26, "H": 46,
          "I": 12, "J": 12, "K": 11, "L": 40, "M": 34, "N": 9, "O": 34}


def build(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Manual Test Cases"

    ws.append(HEADER)
    head_fill = PatternFill("solid", fgColor="D9E1F2")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for row in ROWS:
        ws.append(row)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for col, width in WIDTHS.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30

    wb.save(path)
    return path


if __name__ == "__main__":
    out = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                       "sample_testcases_event_editing.xlsx")
    print("da tao:", build(out))
