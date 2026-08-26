#!/usr/bin/env python3
"""Tests cho 3 mode ghi ket qua (target_resolver) + result_validator.

Diem quan trong nhat duoc kiem chung o day: mode clone_then_write phai giu file
test case GOC khong bi sua mot o nao, va ca run phai ghi vao CUNG MOT ban clone
chu khong clone lai moi case."""
import os
import shutil
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "scripts"))
import result_validator  # noqa: E402
import sheet_io  # noqa: E402
import target_resolver  # noqa: E402

FIXTURE = os.path.join(HERE, "fixtures", "sample_testcases_event_editing.xlsx")
STAMP = "20260824_150000"


class TargetModeTests(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.mkdtemp(prefix="target_mode_")
        self.source = os.path.join(self.tmpdir, "testcases.xlsx")
        shutil.copy2(FIXTURE, self.source)
        self.cfg = {
            "run": {"module_name": "event_editing"},
            "paths": {"execution_runtime_root": os.path.join(self.tmpdir, "runtime")},
            "testcase_source": {"source_type": "excel", "local_excel_path": self.source,
                                "worksheet_name": "Manual Test Cases",
                                "header_row_index": 1, "data_start_row_index": 2},
            "update_policy": {"mode": "in_place"},
        }

    def tearDown(self):
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def _resolve(self):
        return target_resolver.resolve(self.cfg, "event_editing", STAMP)

    def test_target_in_place_writes_into_the_source_file_itself(self):
        # Arrange
        self.cfg["update_policy"]["mode"] = "in_place"
        # Act
        state = self._resolve()
        # Assert
        self.assertEqual(state["target"], self.source)
        self.assertFalse(state["source_untouched"])

    def test_target_backup_then_update_keeps_source_as_target_and_creates_backup(self):
        # Arrange
        self.cfg["update_policy"]["mode"] = "backup_then_update"
        # Act
        state = self._resolve()
        # Assert -- van ghi vao file goc, nhung co ban backup ben canh
        self.assertEqual(state["target"], self.source)
        self.assertTrue(os.path.exists(state["backup"]))
        self.assertNotEqual(state["backup"], self.source)

    def test_target_clone_then_write_creates_clone_and_leaves_source_untouched(self):
        # Arrange
        self.cfg["update_policy"]["mode"] = "clone_then_write"
        # Act
        state = self._resolve()
        # Assert
        self.assertNotEqual(state["target"], self.source)
        self.assertTrue(os.path.exists(state["target"]))
        self.assertTrue(state["source_untouched"])

    def test_target_clone_preserves_every_original_column_value(self):
        # Arrange -- ban clone phai la ban sao dung cau truc, khong phai file tu dinh nghia
        self.cfg["update_policy"]["mode"] = "clone_then_write"
        state = self._resolve()
        from openpyxl import load_workbook
        src_ws = load_workbook(self.source)["Manual Test Cases"]
        clone_ws = load_workbook(state["target"])["Manual Test Cases"]
        # Act + Assert
        self.assertEqual(src_ws.max_row, clone_ws.max_row)
        self.assertEqual(src_ws.max_column, clone_ws.max_column)
        for row in range(1, src_ws.max_row + 1):
            for col in range(1, src_ws.max_column + 1):
                self.assertEqual(src_ws.cell(row=row, column=col).value,
                                 clone_ws.cell(row=row, column=col).value)

    def test_target_clone_run_reuses_same_clone_for_every_case(self):
        # Arrange -- neu khong nho target, moi case se sinh 1 file clone rieng
        self.cfg["update_policy"]["mode"] = "clone_then_write"
        # Act
        first = self._resolve()
        second = target_resolver.resolve(self.cfg, "event_editing", "20260824_160000")
        # Assert
        self.assertEqual(first["target"], second["target"])

    def test_target_reset_starts_a_new_clone_for_the_next_run(self):
        # Arrange
        self.cfg["update_policy"]["mode"] = "clone_then_write"
        first = self._resolve()
        # Act
        target_resolver.reset(self.cfg, "event_editing")
        second = target_resolver.resolve(self.cfg, "event_editing", "20260824_170000")
        # Assert
        self.assertNotEqual(first["target"], second["target"])

    def test_target_invalid_mode_raises(self):
        # Arrange
        self.cfg["update_policy"]["mode"] = "make_up_a_new_file"
        # Act + Assert
        with self.assertRaises(target_resolver.TargetError):
            self._resolve()

    def test_target_clone_write_lands_in_clone_and_source_stays_original(self):
        # Arrange -- end-to-end: ghi that qua sheet_io vao ban clone
        self.cfg["update_policy"].update({
            "mode": "clone_then_write",
            "execution_columns": ["Test result", "Test date", "Tested by", "Remark"],
            "definition_columns": ["TC ID", "Title", "Expected Result"],
        })
        state = self._resolve()
        cfg_target = target_resolver.apply_target(self.cfg, state)
        ws = sheet_io.open_worksheet(cfg_target)
        header = ws.row_values(1)
        mapping = {"test_result": header.index("Test result"),
                   "tc_id": header.index("TC ID")}
        # Act
        sheet_io.write_result(ws, cfg_target, mapping, header, 2, {"test_result": "Pass"})
        # Assert
        from openpyxl import load_workbook
        clone_val = load_workbook(state["target"])["Manual Test Cases"].cell(row=2, column=9).value
        source_val = load_workbook(self.source)["Manual Test Cases"].cell(row=2, column=9).value
        self.assertEqual(clone_val, "Pass")
        self.assertEqual(source_val, "Untested")     # file goc nguyen ven


class ResultValidatorTests(unittest.TestCase):
    CFG = {"result_policy": {"allowed_statuses": ["Pass", "Fail", "N/A", "Untested", "Pending"],
                             "mandatory_remark_when_not_testable": True,
                             "mandatory_remark_when_fail": True}}

    def test_validator_accepts_pass_without_remark(self):
        self.assertTrue(result_validator.validate("Pass", "", "", {}, self.CFG))

    def test_validator_rejects_unknown_status(self):
        with self.assertRaises(result_validator.ResultValidationError):
            result_validator.validate("Blocked", "x", "", {}, self.CFG)

    def test_validator_rejects_pending_without_remark(self):
        with self.assertRaises(result_validator.ResultValidationError):
            result_validator.validate("Pending", "", "", {}, self.CFG)

    def test_validator_rejects_fail_without_actual_when_column_exists(self):
        with self.assertRaises(result_validator.ResultValidationError) as ctx:
            result_validator.validate("Fail", "Actual khác Expected: x", "",
                                      {"actual_result": 12}, self.CFG)
        self.assertIn("Actual Result", str(ctx.exception))

    def test_validator_accepts_fail_without_actual_when_column_absent(self):
        # Template khong co cot Actual Result thi khong the doi
        self.assertTrue(result_validator.validate(
            "Fail", "Actual khác Expected: x", "", {}, self.CFG))


if __name__ == "__main__":
    unittest.main(verbosity=2)
