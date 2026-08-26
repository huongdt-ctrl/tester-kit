#!/usr/bin/env python3
"""
Build a styled Excel test-activity schedule from structured row data.

Usage:
    python3 build_schedule.py <rows.json> <output.xlsx>

rows.json shape:
{
  "project_name": "Optional project/title text",
  "rows": [
    {
      "module": "Login",
      "phase": "1. Phan tich yeu cau & Confirm Q&A",
      "tester": "Anh",
      "effort_hours": 4,               // current estimate, shown as "Est (h)"
      "start_date": "2026-08-24",      // planned
      "end_date": "2026-08-25",        // planned
      "actual_start": null,            // the tester fills these in, in Excel
      "actual_end": null,
      "actual_hours": null,
      "pct_done": null,                // 0-100
      "status": "TODO",
      "note": ""
    },
    ...
  ]
}

The sheet keeps PLAN and ACTUAL side by side on purpose: `Variance (h)` =
actual - estimate is the signal that lets estimate-test recalibrate its base
rates, so overwriting planned dates with real ones would throw away the number
we need. Testers type into the four ACTUAL columns (green header group) and
leave the plan columns alone; sync_schedule.py carries those entries forward on
the next rebuild.

Column layout lives in schedule_format.py. This script only handles formatting,
so every schedule the skill produces looks the same and is safe to recalculate.
"""
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from schedule_format import (ACTUAL_HEADER_FILL, CENTERED, COLUMNS, DATA_START, DATE_COLS,
                             DURATION_COL, FONT_NAME, HEADER_FONT_COLOR, HEADER_ROW, INPUT_FILL,
                             LAST_COL, PLAN_HEADER_FILL, STATUS_FILL, STATUS_OPTIONS,
                             VARIANCE_COL, parse_date)

THIN = Side(style="thin", color="B7B7B7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def write_title(ws, project_name):
    ws.merge_cells(f"A1:{LAST_COL}1")
    cell = ws["A1"]
    cell.value = project_name
    cell.font = Font(name=FONT_NAME, size=14, bold=True)
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 26


def write_header(ws):
    for col_idx, (_key, header, width, group) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=HEADER_ROW, column=col_idx, value=header)
        fill = ACTUAL_HEADER_FILL if group == "actual" else PLAN_HEADER_FILL
        cell.font = Font(name=FONT_NAME, size=11, bold=True, color=HEADER_FONT_COLOR)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    # keep No/Module/Phase/Tester visible while scrolling right to the actuals
    ws.freeze_panes = "E3"


def write_row(ws, r, index, row):
    status = str(row.get("status", "TODO")).upper()
    if status not in STATUS_OPTIONS:
        status = "TODO"

    for col_idx, (key, _header, _w, group) in enumerate(COLUMNS, start=1):
        if key == "no":
            value = index + 1
        elif key is None:
            value = None  # formula, filled in below
        elif col_idx in DATE_COLS:
            value = parse_date(row.get(key))
        elif key == "status":
            value = status
        else:
            value = row.get(key, "")

        cell = ws.cell(row=r, column=col_idx, value=value)
        cell.font = Font(name=FONT_NAME, size=11)
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center" if col_idx in CENTERED else None,
                                   vertical="center", wrap_text=True)
        if group == "actual":
            cell.fill = PatternFill("solid", fgColor=INPUT_FILL)
        if col_idx in DATE_COLS:
            cell.number_format = "yyyy-mm-dd"

    ws.cell(row=r, column=DURATION_COL).value = f"=G{r}-F{r}+1"
    ws.cell(row=r, column=11).number_format = "0.00"   # Actual (h)
    ws.cell(row=r, column=12).number_format = "0"      # % Done
    # blank until Actual (h) is filled in, so an untouched row doesn't read as -Est
    variance = ws.cell(row=r, column=VARIANCE_COL)
    variance.value = f'=IF(K{r}="","",K{r}-E{r})'
    variance.number_format = "0.00"

    ws.cell(row=r, column=14).fill = PatternFill("solid", fgColor=STATUS_FILL.get(status, "FFFFFF"))


def add_validation(ws, last_row):
    status_dv = DataValidation(type="list", allow_blank=True,
                               formula1='"{}"'.format(",".join(STATUS_OPTIONS)))
    ws.add_data_validation(status_dv)
    status_dv.add(f"N{DATA_START}:N{last_row}")

    pct_dv = DataValidation(type="whole", operator="between", formula1=0, formula2=100,
                            allow_blank=True, errorTitle="Giá trị không hợp lệ",
                            error="% Done phải là số nguyên 0-100.")
    ws.add_data_validation(pct_dv)
    pct_dv.add(f"L{DATA_START}:L{last_row}")

    ws.auto_filter.ref = f"A{HEADER_ROW}:{LAST_COL}{last_row}"


def build(rows_path, output_path):
    with open(rows_path, "r", encoding="utf-8") as f:
        data = json.load(f)
    rows = data["rows"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Test Schedule"

    write_title(ws, data.get("project_name", "Test Schedule"))
    write_header(ws)
    for i, row in enumerate(rows):
        write_row(ws, DATA_START + i, i, row)

    last_row = DATA_START + len(rows) - 1
    if last_row >= DATA_START:
        add_validation(ws, last_row)

    wb.save(output_path)
    print(json.dumps({"status": "ok", "rows_written": len(rows), "output": output_path}))


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: python3 build_schedule.py <rows.json> <output.xlsx>", file=sys.stderr)
        sys.exit(1)
    build(sys.argv[1], sys.argv[2])
