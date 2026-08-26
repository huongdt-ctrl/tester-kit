#!/usr/bin/env python3
"""
Fill the cached values of the `Duration (days)` formula column in a schedule
built by build_schedule.py.

Why this exists: openpyxl writes formulas with an empty <v> cache, so anything
that reads the file WITHOUT a spreadsheet engine (pandas, markitdown, previews)
sees Duration as blank. Excel / Google Sheets recalculate on open anyway --
build_schedule.py leaves fullCalcOnLoad="1" set -- so this only repairs the
cache for programmatic readers. The formula itself is left untouched: the value
is never hardcoded in place of `=G{r}-F{r}+1`.

Use this in place of the platform recalc.py when that script (or LibreOffice)
is not available.

Usage:
    python3 recalc_duration.py <schedule.xlsx>
"""
import json
import re
import shutil
import sys
import zipfile

from openpyxl import load_workbook

DURATION_COL = 8  # H -- keep in sync with COLUMNS in build_schedule.py
DATA_START = 3    # row 1 = title, row 2 = header


def duration_values(path):
    """Compute end-start+1 per data row, mirroring the sheet formula."""
    ws = load_workbook(path).active
    values = {}
    errors = []
    for r in range(DATA_START, ws.max_row + 1):
        start = ws.cell(row=r, column=6).value
        end = ws.cell(row=r, column=7).value
        if start is None or end is None:
            continue
        try:
            days = (end - start).days + 1
        except TypeError:
            errors.append(f"row {r}: start/end is not a date ({start!r}, {end!r})")
            continue
        if days < 1:
            errors.append(f"row {r}: end date is before start date")
            continue
        values[f"H{r}"] = days
    return values, errors


def inject(path, values):
    """Rewrite sheet1.xml with <v> caches filled for the duration cells."""
    with zipfile.ZipFile(path) as zin:
        members = {n: zin.read(n) for n in zin.namelist()}

    sheet = members["xl/worksheets/sheet1.xml"].decode("utf-8")
    patched = 0
    for ref, val in values.items():
        pattern = re.compile(r'(<c r="%s"[^>]*>)(<f>[^<]*</f>)<v>\s*</v>' % ref)
        sheet, n = pattern.subn(lambda m: f"{m.group(1)}{m.group(2)}<v>{val}</v>", sheet)
        patched += n
    members["xl/worksheets/sheet1.xml"] = sheet.encode("utf-8")

    shutil.copy(path, path + ".bak")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in members.items():
            zout.writestr(name, data)
    return patched


def main():
    if len(sys.argv) != 2:
        print("Usage: python3 recalc_duration.py <schedule.xlsx>", file=sys.stderr)
        sys.exit(1)
    path = sys.argv[1]
    values, errors = duration_values(path)
    patched = inject(path, values) if values else 0
    result = {
        "status": "success" if not errors and patched == len(values) else "error",
        "cells_patched": patched,
        "cells_expected": len(values),
        "total_errors": len(errors),
    }
    if errors:
        result["errors"] = errors
    print(json.dumps(result, ensure_ascii=False))
    sys.exit(0 if result["status"] == "success" else 1)


if __name__ == "__main__":
    main()
