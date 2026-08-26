#!/usr/bin/env python3
"""
Single source of truth for the schedule sheet's layout.

build_schedule.py writes these columns; sync_schedule.py reads them back out of
an existing file. Both import from here so a column can be added or moved in one
place without the reader and the writer drifting apart.
"""
from datetime import date, datetime

from openpyxl.utils import get_column_letter

STATUS_OPTIONS = ["TODO", "DOING", "DONE", "PENDING", "CANCEL"]
STATUS_FILL = {
    "TODO": "D9D9D9",     # gray
    "DOING": "FFF2CC",    # light amber
    "DONE": "C6EFCE",     # light green
    "PENDING": "FCE4D6",  # light orange
    "CANCEL": "F4CCCC",   # light red
}

# Statuses whose plan dates are frozen -- the work is finished, so re-planning it
# would rewrite history. See sync_schedule.py.
FROZEN_STATUSES = {"DONE", "CANCEL"}

PLAN_HEADER_FILL = "305496"    # dark blue -- planned / identity columns
ACTUAL_HEADER_FILL = "1F7246"  # dark green -- columns the tester fills in
INPUT_FILL = "F2F9F4"          # faint green wash over the ACTUAL data cells
HEADER_FONT_COLOR = "FFFFFF"
FONT_NAME = "Calibri"

# (key, header, width, group). key = the field name in rows.json;
# None means the cell holds a formula rather than caller-supplied data.
COLUMNS = [
    ("no",           "No",               5,  "plan"),
    ("module",       "Module / Feature", 26, "plan"),
    ("phase",        "Test Phase",       30, "plan"),
    ("tester",       "Tester",           13, "plan"),
    ("effort_hours", "Est (h)",           9, "plan"),
    ("start_date",   "Plan Start",       12, "plan"),
    ("end_date",     "Plan End",         12, "plan"),
    (None,           "Duration (days)",  11, "plan"),
    ("actual_start", "Actual Start",     12, "actual"),
    ("actual_end",   "Actual End",       12, "actual"),
    ("actual_hours", "Actual (h)",       10, "actual"),
    ("pct_done",     "% Done",            8, "actual"),
    (None,           "Variance (h)",     11, "actual"),
    ("status",       "Status",           11, "plan"),
    ("note",         "Note",             36, "plan"),
]

HEADER_ROW = 2
DATA_START = 3

# 1-based column index per rows.json key, e.g. COL["actual_hours"] == 11
COL = {key: i for i, (key, *_) in enumerate(COLUMNS, start=1) if key}
DURATION_COL = 8
VARIANCE_COL = 13
LAST_COL = get_column_letter(len(COLUMNS))

CENTERED = {1, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14}
DATE_COLS = {COL["start_date"], COL["end_date"], COL["actual_start"], COL["actual_end"]}


def parse_date(value):
    """Accept an ISO string, a date/datetime, or blank -> None."""
    if value in (None, ""):
        return None
    if isinstance(value, (date, datetime)):
        return value
    return datetime.strptime(str(value)[:10], "%Y-%m-%d").date()


def iso(value):
    """Inverse of parse_date: date/datetime -> 'YYYY-MM-DD', blank -> None."""
    d = parse_date(value)
    if d is None:
        return None
    return (d.date() if isinstance(d, datetime) else d).isoformat()


# Header text -> rows.json key, used when READING an existing file. Resolving by
# header rather than by fixed index lets sync_schedule.py ingest schedules built
# by an older layout (v1 had 10 columns and called Est "Effort (h)") without
# silently reading Status as a date.
HEADER_ALIASES = {
    "Effort (h)": "effort_hours",
    "Start Date": "start_date",
    "End Date": "end_date",
}

HEADER_TO_KEY = {header: key for key, header, *_ in COLUMNS if key}
HEADER_TO_KEY.update(HEADER_ALIASES)


def resolve_columns(ws):
    """{rows.json key: 1-based column index} read off the sheet's header row."""
    found = {}
    for cell in ws[HEADER_ROW]:
        key = HEADER_TO_KEY.get(str(cell.value).strip() if cell.value else "")
        if key:
            found[key] = cell.column
    return found
